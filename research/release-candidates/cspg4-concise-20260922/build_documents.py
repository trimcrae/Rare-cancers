from pathlib import Path
import re,json,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
W=Path(__file__).resolve().parent;O=W/'journal-package';O.mkdir(exist_ok=True)
def inline(p,s):
    for part in re.split(r'(\*\*.*?\*\*)',s):
        r=p.add_run(part.strip('*') if part.startswith('**') else part);r.bold=part.startswith('**')
def setup(doc,supp=False):
    # The bundled base template carries a decorative title border. Remove it explicitly.
    for root in [doc.styles.element,doc._element]:
        for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
    s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.75);s.left_margin=s.right_margin=Inches(.85)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11 if supp else 12)
    normal.paragraph_format.line_spacing=1.12 if supp else 1.4;normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',17),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.font.bold=True
        st.paragraph_format.space_before=Pt(12 if name!='Title' else 0);st.paragraph_format.space_after=Pt(7);st.paragraph_format.keep_with_next=True
    f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.paragraph_format.space_after=Pt(0)
    r=f.add_run('CSPG4 RNA across sarcoma types  |  ');r.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    doc.core_properties.author='Tristan D. McRae';doc.core_properties.subject='Comparative sarcoma RNA study'
def table(doc,rows):
    rows=[[v.strip() for v in r.strip('|').split('|')] for r in rows if not re.fullmatch(r'[\s|:\-]+',r)]
    t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if len(rows[0])==6: widths=[2.53,1.08,.36,.72,1.23,.58]
    elif len(rows[0])==4 and rows[0][0]=='Gene':widths=[1.0,1.5,2.0,2.0]
    elif len(rows[0])==4 and rows[0][1]=='Hofvander S1 row':widths=[.8,.9,2.35,2.75]
    elif len(rows[0])==4 and rows[0][0]=='Specimen ID':widths=[.90,1.45,.55,3.40]
    elif len(rows[0])==4:widths=[1.02,1.65,1.10,2.9]
    else:widths=[1.50,2.0,3.27]
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for i,vals in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(c,v) in enumerate(zip(cells,vals)):
            if len(rows[0])==6 and j==1:v=v.replace(';','; ')
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            if len(rows[0])==6 and j>=2 or rows[0][0]=='Gene' and j>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            rr=p.add_run(v);rr.font.size=Pt(9.5 if len(rows[0])==6 else 10);rr.bold=i==0
            tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
            tcpr.append(b);m=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');m.append(el)
            tcpr.append(m)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5ECF0');tcpr.append(sh)
        pr=t.rows[i]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
        if i==0:repeat=OxmlElement('w:tblHeader');pr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
def prose_content(content):
    if content.startswith('---\n'):
        end=content.find('\n---\n',4)
        if end<0:raise ValueError('Unclosed repository metadata')
        return content[end+5:]
    return content

def markdown(doc,content,supp=False):
    content=prose_content(content)
    lines=content.splitlines();i=0;abstract=False;references=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('|'):
            block=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
            table(doc,block);continue
        if line.startswith('# '):doc.add_paragraph(line[2:],style='Title');continue
        if line.startswith('## '):
            title=line[3:]
            if title=='Introduction':doc.add_page_break()
            abstract=title=='Abstract';references=title in ['References','Supplementary references']
            doc.add_heading(title,level=1);continue
        if line.startswith('### '):doc.add_heading(line[4:],level=2);continue
        p=doc.add_paragraph();inline(p,line)
        if abstract:
            p.paragraph_format.line_spacing=1.1
            for r in p.runs:r.font.size=Pt(11)
        if references:
            p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(7)
            for r in p.runs:r.font.size=Pt(10.5)
def figure(doc,n,caption):
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11);sec.page_height=Inches(8.5)
    sec.top_margin=sec.bottom_margin=Inches(.45);sec.left_margin=sec.right_margin=Inches(.65)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1
    p.add_run().add_picture(str(W/'figures'/f'{n}.png'),width=Inches(9.5))
    p=doc.add_paragraph();inline(p,caption);p.paragraph_format.line_spacing=1.04;p.paragraph_format.space_after=Pt(0)
    for r in p.runs:r.font.size=Pt(9)
main=(W/'manuscript.md').read_text(encoding='utf8');body,leg=main.split('## Figure legends\n',1)
caps=[s.strip() for s in leg.strip().split('\n\n')]
doc=Document();setup(doc);markdown(doc,body)
for n,c in zip(['figure-1','figure-2'],caps):figure(doc,n,c)
doc.save(O/'CSPG4-manuscript.docx')
supp=(W/'supplement.md').read_text(encoding='utf8')
doc=Document();setup(doc,True);markdown(doc,supp,True)
figure(doc,'figure-S1','**Figure S1. Other eligible tumors in the Hofvander source.** The 242 specimens outside the defined malignant comparison span 31 source labels. Dots show specimens, segments show interquartile ranges, and vertical marks show medians. The axis uses log2(1+TPM) spacing with TPM labels. GIST, gastrointestinal stromal tumor; NOS, not otherwise specified.')
figure(doc,'figure-S2','**Figure S2. External within-source CSPG4 ordering against LMS.** Each cell compares the row diagnosis with LMS within its source and requires at least five profiles per group; blanks lack that support. B denotes published Boudin cohorts; TH denotes Treehouse original accessions. Each retains its measurement scale and historical mixed-platform qualifications.')

doc.save(O/'CSPG4-supplement.docx')

(O/'figure-legends.txt').write_text(leg.strip()+'\n',encoding='utf8')
(W/'DOCUMENT-BUILD.json').write_text(json.dumps({'source_hashes':{n:hashlib.sha256((W/n).read_bytes()).hexdigest() for n in ['manuscript.md','supplement.md','build_documents.py']},'docx':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in O.glob('*.docx')},'status':'awaiting actual remote render and every-page visual/text verification'},indent=2)+'\n')
print(json.dumps({'docx':[p.name for p in O.glob('*.docx')]}))
