"""Create editable Cancer Genetics files from the exact submitted ATR v3 text."""
from pathlib import Path
import hashlib, importlib.util, json, re, shutil
from lxml import html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
BASE = Path('C:/Users/mcrae/.codex/review-workspaces/opus-capacity-sprint-20260907/capacity-campaign-20260908/observations/atr-preprint-release-20260909/frozen-candidate-v3')
TITLE = 'Reading-frame constraints and retained RG content in NR4A3 fusion models of extraskeletal myxoid chondrosarcoma'
ABSTRACT = '''**Background:** NR4A3 fusions define extraskeletal myxoid chondrosarcoma (EMC), but their recruitment to DNA double-strand breaks and response to ATR inhibition remain untested. We examined reference-model reading frames and retained FET-partner RG content to delimit sequence-based predictions.

**Methods:** Reported exon junctions from primary studies and reviews were assembled from public reference transcripts and translated at transcript level. Retained RG dipeptides were compared with published recruitment anchors. TCF12 composition was compared with FET proteins on a symmetric prefix grid.

**Results:** Four sourced junctions produced in-frame models retaining the complete NR4A3 moiety. The EWSR1 exon-7/NR4A3 exon-2 model encoded a 59-residue insertion absent from the earlier protein-level model. EWSR1 type-1 and type-2 models retained 8/30 and 0/30 RG dipeptides, respectively, within the 0.000–0.267 span of three reported EWSR1::ATF1 breakpoints. Those breakpoint models are distinct from the measured construct, whose breakpoint is unstated; measured anchors are 0.000 and 1.000. TCF12 remained outside the FET compositional range across the evaluated grid of prefixes from 50 residues in ten-residue steps.

**Conclusions:** Transcript-level assembly resolves model-specific reading-frame constraints and supports bounded, falsifiable predictions. Single-source reference annotation and unverified patient junctions limit interpretation. No experiment, recruitment measurement or ATR-inhibitor-response test was performed.'''
KEYWORDS = 'extraskeletal myxoid chondrosarcoma; NR4A3 fusion; reading frame; EWSR1; TCF12; RG dipeptides'
HIGHLIGHTS = [
    'Four sourced junctions produce in-frame NR4A3 reference fusion models.',
    'Transcript-level assembly reveals a 59-residue type-2 insertion.',
    'Retained RG content supports bounded predictions, not measured recruitment.',
    'TCF12 lies outside the FET compositional range on the evaluated grid.',
    'Patient junctions and ATR-inhibitor response remain unvalidated.',
]

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def base_doc(title):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(.85)
    for name, size in [('Normal', 11), ('Title', 16), ('Heading 1', 13), ('Heading 2', 12), ('Heading 3', 11)]:
        st = d.styles[name]
        st.font.name = 'Times New Roman'
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.5
    for st in d.styles:
        for el in list(st.element.xpath('.//w:pBdr')):
            el.getparent().remove(el)
        for el in st.element.xpath('.//w:rFonts'):
            for a in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme', 'csTheme']:
                el.attrib.pop(qn('w:'+a), None)
    footer = s.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    d.core_properties.author = 'Tristan D. McRae'
    d.core_properties.title = title
    return d

def inline(p, node, bold=False, italic=False, super_=False, sub=False, mono=False):
    def add(text):
        if not text:
            return
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        r.font.superscript, r.font.subscript = super_, sub
        if mono:
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
    add(node.text)
    for child in node:
        tag = child.tag.lower() if isinstance(child.tag, str) else ''
        if tag == 'br':
            p.add_run().add_break()
        elif tag == 'a' and child.get('href', '').startswith('http'):
            link = OxmlElement('w:hyperlink')
            rid = p.part.relate_to(child.get('href'), 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
            link.set(qn('r:id'), rid)
            rr = OxmlElement('w:r'); tt = OxmlElement('w:t')
            tt.text = ''.join(child.itertext()); rr.append(tt); link.append(rr); p._p.append(link)
        else:
            inline(p, child, bold or tag in ['b', 'strong'], italic or tag in ['i', 'em'], super_ or tag=='sup', sub or tag=='sub', mono or tag=='code')
        add(child.tail)

def make_doc(md, title, renderer):
    d = base_doc(title)
    # Preserve source blockquote text as ordinary paragraphs and suppress only the
    # standalone image placeholder; the original image is embedded separately.
    md = re.sub(r'^> ?', '', md, flags=re.M)
    md = re.sub(r'^!\[.*?\]\([^\n]+\)\s*$', '', md, flags=re.M)
    tree = html.fragment_fromstring(renderer.markdown_to_html(md), create_parent='div')
    def walk(node):
        tag = node.tag.lower() if isinstance(node.tag, str) else ''
        if tag in ['ol', 'ul']:
            start = int(node.get('start', '1'))
            for i, child in enumerate(node):
                p = d.add_paragraph()
                p.add_run(str(start+i)+'. ' if tag=='ol' else '\u2022 ')
                inline(p, child)
            return
        if tag in ['p', 'h1', 'h2', 'h3', 'h4', 'pre', 'li']:
            style = {'h1':'Title', 'h2':'Heading 1', 'h3':'Heading 2', 'h4':'Heading 3'}.get(tag)
            p = d.add_paragraph(style=style)
            inline(p, node, mono=tag=='pre')
        elif tag == 'table':
            rows = node.xpath('.//tr')
            count = max(len(r.xpath('./th|./td')) for r in rows)
            table = d.add_table(rows=0, cols=count)
            table.autofit = False
            widths = [1.6,5.2] if count==2 else [6.8/count]*count
            for col,w in zip(table.columns,widths): col.width=Inches(w)
            for i,row in enumerate(rows):
                cells=table.add_row().cells
                parts=row.xpath('./th|./td')
                if len(parts)==1 and int(parts[0].get('colspan','1'))==count:
                    cells=[cells[0].merge(cells[-1])]
                no_split=OxmlElement('w:cantSplit')
                table.rows[-1]._tr.get_or_add_trPr().append(no_split)
                for cell,content,w in zip(cells,row.xpath('./th|./td'),widths):
                    cell.width=Inches(6.8 if len(cells)==1 else w); p=cell.paragraphs[0]
                    p.paragraph_format.line_spacing=1.0
                    p.paragraph_format.space_after=Pt(4)
                    p.paragraph_format.space_before=Pt(4)
                    inline(p,content,bold=content.tag=='th')
                    for run in p.runs:run.font.size=Pt(9)
                    pr=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
                    for edge in ['top','bottom','left','right']:
                        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                    pr.append(borders)
                if all(c.tag=='th' for c in parts):
                    mark=OxmlElement('w:tblHeader');table.rows[-1]._tr.get_or_add_trPr().append(mark)
            d.add_paragraph()
        elif tag not in ['hr', 'img']:
            for child in node:walk(child)
    for node in tree:walk(node)
    return d

def build():
    assert sha(BASE/'emc-atr-collaborator-package.pdf')=='cd87f56acda32619619444e672974a0faec4cf016b05cfe3cbe0c272dfd722fb'
    source=(BASE/'emc-atr-collaborator-package.md').read_text(encoding='utf8')
    (HERE/'baseline-source.md.original.txt').write_text(source,encoding='utf8')
    clean=re.sub(r'^---\n.*?\n---\n','',source,flags=re.S)
    clean=re.sub(r'<!--.*?-->','',clean,flags=re.S)
    clean=re.sub(r'\*\*Keywords:\*\*.*?\n\n','**Keywords:** '+KEYWORDS+'\n\n',clean,count=1,flags=re.S)
    clean=re.sub(r'## Abstract\s+.*?(?=\n---\s*\n\s*## 1\.)','## Abstract\n\n'+ABSTRACT+'\n',clean,count=1,flags=re.S)
    main,supp=clean.split('## 9. Supplementary tables',1)
    main=main.replace('## 4. Pre-specified predictions','## 4. Discussion and pre-specified predictions')
    main=main.replace('Two drawing-level\nobservations, on the legibility of the closely spaced NR4A3 ticks and on the binding between the Panel A\nbox literals and their source data field, are recorded in the integration\nQA note cited in the editorial comment above rather than in this caption.','')
    main=re.sub(r'Two drawing-level\s+observations,.*?rather than in this caption\.', '', main, flags=re.S)
    old_status=re.search(r'\*\*The registered text above.*?now known about it\.',main,re.S)
    if not old_status:raise ValueError('Prediction presentation status anchor missing')
    main=main.replace(old_status.group(0),'**Table 5 presents the manuscript prediction set alongside the pinned prediction records.**\nThe immutable records and this manuscript presentation are distinct records, as detailed below.\nCorrections and limitations are stated separately; neither original prediction record is rewritten here.')
    reproduction='''### 2.5 Reproduction

From this journal-package repository revision, run:

```
cd research/release-candidates/PUB-ATR-PANEL-ASK
cd 2026-09-10-cancer-genetics-r1
python3 verify_study_scope.py
```

This read-only wrapper verifies the original six-model scope (EWSR1, TAF15, FUS, TCF12,
TFG and NR4A3), compares every field of the derived construct artifact with the stored artifact,
runs the unchanged frame-and-composition checker, and checks the original figure's packaged
provenance stamp. It prints `ORIGINAL STUDY SCOPE REPRODUCES`, `REPRODUCES`, and
`PROVENANCE MATCHES` only when those respective checks pass. Gene-model fields, UniProt
sequences and pinned dependencies are hash-checked; a changed study input fails verification.
No cache, producer, prediction record, scientific artifact or figure is rewritten.

The unscoped construct producer's `--check` reports metadata drift in `gene_models` and
`ensembl_vs_uniprot_sequences` because the shared input cache also contains PGR, which was
outside the original study model set. The wrapper explicitly excludes that extra model in
memory; it does not discard fields or differences within the study models. The entire original-scope
derived artifact matches. The figure check explicitly uses this package's original provenance
record rather than a different stamp elsewhere in the checkout. These are verification checks
of retained results, not a new scientific analysis.

'''
    main,n=re.subn(r'### 2\.5 Reproduction\s+.*?(?=Retrieval, computation and drafting)',reproduction,main,count=1,flags=re.S)
    if n!=1:raise ValueError('Reproduction section anchor missing')
    main=main.replace("and its producer's `--check` re-derives it offline.","and the original-study-scope wrapper in section 2.5 re-derives it offline.")
    declarations='''## Author declarations

**Funding:** No funding was received.

**Competing interests:** The author declares no competing interests.

**Author contribution:** Tristan D. McRae directed the work and is responsible for its content, as described in section 2.5.

## Declaration of generative AI and AI assisted technologies

Substantial AI assistance in retrieval, computation and drafting is disclosed in section 2.5. OpenAI Codex additionally assisted with this journal presentation and file verification. The author directed the work and is responsible for its content.

'''
    main=main.replace('## 8. References',declarations+'## 8. References')
    references=main.split('## 8. References',1)[1]
    si='# Supplementary information\n\n'+TITLE+'\n\nThese tables retain the submitted Research Square v3 scientific content. References use the numbering in the main article.\n\n## Supplementary tables\n'+supp+'\n\n## References\n'+references
    for name,text in [('manuscript-source.txt',main),('supplement-source.txt',si)]:
        (HERE/name).write_text(text,encoding='utf8')
    spec=importlib.util.spec_from_file_location('bsp',ROOT/'research/manuscripts/build_submission_pdf.py')
    renderer=importlib.util.module_from_spec(spec);spec.loader.exec_module(renderer)
    d=make_doc(main,TITLE,renderer)
    d.add_page_break(); d.add_paragraph('Figure 1','Heading 1')
    d.add_picture(str(BASE/'emc-fusion-frame-fig1.png'),width=Inches(6.8))
    d.save(HERE/'manuscript.docx')
    make_doc(si,'Supplementary information',renderer).save(HERE/'supplementary-information.docx')
    for name in ['emc-fusion-frame-fig1.png','emc-fusion-frame-fig1.pdf','emc-atr-figure-provenance.json']:
        shutil.copyfile(BASE/name,HERE/name)
    legend=re.search(r'\*\*Figure 1\..*?(?=\n\*\*Table 3\.)',main,re.S)
    if not legend:raise ValueError('Original figure legend missing')
    make_doc('# Figure legend\n\n'+legend.group(0),'Figure legend',renderer).save(HERE/'figure-legend.docx')
    assert len(re.sub(r'\*\*','',ABSTRACT).split())<=250
    assert all(len(x)<=85 for x in HIGHLIGHTS)
    make_doc('# Highlights\n\n'+'\n\n'.join(HIGHLIGHTS),'Highlights',renderer).save(HERE/'highlights.docx')
    (HERE/'highlights.txt').write_text('\n'.join(HIGHLIGHTS)+'\n',encoding='utf8')
    cover='''# Cover letter

Dear Editors,

Please consider “'''+TITLE+'''” as an Original Research Article in Cancer Genetics.

This computational study reconstructs four sourced NR4A3 fusion junctions at transcript level, resolves reference-model reading-frame constraints, and places retained RG content on a published axis while distinguishing breakpoint models from the measured construct. It also reports the retained symmetric TCF12 composition comparison. No recruitment experiment or ATR-inhibitor-response test was performed, and patient-junction confirmation remains necessary.

The scientific content is unchanged from the existing Research Square submission rs-10988531, currently in prescreening. This journal package supplies editable text and tables, a structured abstract, separate supplementary tables and original figure files. It does not create a duplicate preprint submission or claim that the preprint is publicly posted.

We request the standard subscription route with online-only color and no optional paid open-access or accelerated service. Funding, interests and AI assistance are disclosed in the manuscript.

Sincerely,

Tristan D. McRae

Independent researcher, unaffiliated

trimcrae@gmail.com
'''
    make_doc(cover,'Cover letter',renderer).save(HERE/'cover-letter.docx')
    manifest={'baseline_pdf_sha256':sha(BASE/'emc-atr-collaborator-package.pdf'),'baseline_source_sha256':sha(BASE/'emc-atr-collaborator-package.md'),'renderer_sha256':sha(ROOT/'research/manuscripts/build_submission_pdf.py'),'generator_sha256':sha(Path(__file__)),'abstract_words':len(re.sub(r'\*\*','',ABSTRACT).split()),'keywords':6,'highlights_characters':[len(x) for x in HIGHLIGHTS],'scientific_analysis_run':False,'render_and_review':'pending','outputs':[{'path':p.name,'bytes':p.stat().st_size,'sha256':sha(p)} for p in sorted(HERE.iterdir()) if p.suffix in ['.docx','.txt','.pdf','.png']]}
    (HERE/'generation-manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf8')
    print(json.dumps({'abstract_words':manifest['abstract_words'],'outputs':len(manifest['outputs']),'tables_main':len(d.tables),'render':'pending'}))

if __name__=='__main__':build()
