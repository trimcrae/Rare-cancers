"""Build author documents and byte-preserving data package; no scientific rerun.

Requires python-docx for documents. Run from any working directory.
PDFs are rendered separately with the canonical DOCX renderer.
"""
from pathlib import Path
import csv
import hashlib
import json
import re
import shutil
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
BASE = ROOT.parent
CAT = BASE / '2026-09-30-full-catalogue'
OLD = BASE / '2026-09-26/evidence'
SHA = 'a916dab2979e27f930b417243f021b6c2b2ca371'

def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def rows(p):
    with p.open(encoding='utf-8', newline='') as f:
        return list(csv.DictReader(f, delimiter='\t'))

def dump(p, obj):
    p.write_text(json.dumps(obj, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')

def body(p):
    return p.read_text(encoding='utf-8').split('\n---\n', 1)[1].strip()

def document(supp=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.75)
    for name, size in [('Normal', 11), ('Title', 16), ('Heading 1', 11)]:
        s = doc.styles[name]
        s.font.name, s.font.size = 'Times New Roman', Pt(size)
        s.font.color.rgb = RGBColor(0, 0, 0)
        for border in list(s.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        fonts = s.element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is not None:
            for key in list(fonts.attrib):
                if key.endswith('Theme'):
                    del fonts.attrib[key]
            for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
                fonts.set(qn('w:' + key), 'Times New Roman')
        s.paragraph_format.space_after = Pt(5)
        s.paragraph_format.line_spacing = 1.05
        s.paragraph_format.widow_control = True
    doc.styles['Heading 1'].font.bold = True
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
    doc.styles['Heading 1'].paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    r = footer.add_run('Supporting material | ' if supp else 'Author proof | ')
    r.font.size = Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.author = 'Tristan D. McRae'
    return doc

def para(doc, text, small=False):
    p = doc.add_paragraph(text)
    if small:
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1
        for r in p.runs:
            r.font.size = Pt(9.5)
    return p

def table(doc, headers, data, widths, size=9.5, bold_rows=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.style = 'Table Grid'
    for c, width in zip(t.columns, widths):
        c.width = Inches(width)
    for row_i, values in enumerate([headers] + data):
        row = t.rows[0] if row_i == 0 else t.add_row()
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if row_i == 0:
            props.append(OxmlElement('w:tblHeader'))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Inches(width)
            cell.text = str(value)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.line_spacing = 1
                for r in p.runs:
                    r.font.size = Pt(size)
                    r.bold = row_i == 0 or row_i - 1 in bold_rows
            if row_i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'EEEEEE')
                cell._tc.get_or_add_tcPr().append(shade)
    return t

designs = rows(CAT / 'results/all-designs.tsv')
junctions = rows(CAT / 'results/junction-catalogue.tsv')
example = sorted([r for r in designs if r['junction'] == 'TCF12_e5__NR4A3_e3'], key=lambda r: int(r['donor_bases']))
assert len(example) == 5 and len(designs) == 220 and len(junctions) == 44
table1 = [[r[k] for k in ['donor_bases', 'canonical_seven_gene_gap_run_bp', 'expanded_seven_gene_gap_run_bp', 'expanded_seven_gene_min_hamming_bp']] for r in example]
table1_caption = 'Table 1. Including normal RNA variants changes the preferred position at the deposited TCF12-NR4A3 junction.'
table1_note = ('Match length is the longest exact normal-RNA match spanning the central six DNA bases; shorter is favored. '
               'The final column is the fewest differences from any complete normal 16-base window; more is favored among ties. '
               'The last column uses the expanded comparison. Upstream bases come from TCF12. '
               'All comparisons use the same seven genes. Bold marks the final choice: ASO 5\'-GGCATATCCATCAGAT-3\'.')
table1_md = table1_caption + '\n\n| Upstream bases | One RNA per gene: match (bases) | With RNA variants: match (bases) | Fewest differences: expanded |\n|---|---|---|---|\n' + '\n'.join('| ' + ' | '.join(r) + ' |' for r in table1) + '\n\n' + table1_note
(ROOT / 'table1.txt').write_text(table1_md + '\n', encoding='utf-8')

def add_table1(doc):
    para(doc, table1_caption, True).paragraph_format.keep_with_next = True
    table(doc, ['Upstream\nbases', 'One RNA per gene:\nmatch (bases)', 'With RNA variants:\nmatch (bases)', 'Fewest differences\n(expanded comparison)'], table1, [.85, 1.85, 1.85, 2.45], bold_rows=[3])
    para(doc, table1_note, True)

def add_evidence(doc):
    para(doc, 'Table S1. Source-supported distinct reference junctions.', True).paragraph_format.keep_with_next = True
    data = [
        ['TAF15 e6 / NR4A3 e3', 'Deposit', 'AF162670.1; AJ243810.1; AJ245932.1'],
        ['TFG e7 / NR4A3 e3', 'Deposit', 'AY532911.1'],
        ['TCF12 e5 / NR4A3 e3', 'Deposit', 'AF289510.1'],
        ['EWSR1 e10 / NR4A3 cryptic', 'Deposit', 'AF524261.1'],
        ['EWSR1 e7 / NR4A3 e2', 'Deposit', 'S81242.1 (transcribed from print)'],
        ['EWSR1 e12 / NR4A3 e3', 'Reconstruction', 'USZ20 coordinates; separate Brenca E-N description'],
        ['TAF15 e6 / NR4A3 cryptic', 'Reconstruction', 'Brenca T-N construct description'],
    ]
    table(doc, ['Reference junction', 'Evidence', 'Source'], data, [2.2, 1.2, 3.6])
    para(doc, 'Labels use this catalogue\'s reference convention. Deposit support is local junction correspondence, not proof of a complete patient transcript or an independent-patient count. Data S7 supplies breakpoints and ambiguity.', True)

def add_catalogue(doc):
    doc.add_page_break()
    doc.add_paragraph('Table S2. Complete catalogue of final sequence choices', 'Heading 1')
    para(doc, 'D = deposited local junction; R = reference reconstruction; H = hypothetical exact join; C = annotation-error control. Each row retains all final ties. Labels are reference conventions. Gap = longest gap-spanning match; diff = minimum whole-window differences. Values refer to the expanded seven-gene corpus. The control is excluded from target totals.', True)
    data = []
    codes = {'deposited_sequence_match':'D', 'coordinate_based_reference_reconstruction':'R', 'construct_description_reference_reconstruction':'R', 'hypothetical_exact_reference_join':'H', 'annotation_error_control':'C'}
    for r in junctions:
        label = r['junction'].replace('__NR4A3_', ' / ').replace('_', ' ').replace('intron2crypticExon','cryptic')
        positions = json.loads(r['final_co_winner_donor_bases'])
        seqs = json.loads(r['final_co_winner_ASO_5to3'])
        data.append([label, codes[r['evidence_class']], '\n'.join(f'{p}: {s}' for p, s in zip(positions,seqs)), f"{r['primary_best_gap_run_bp']} / {r['final_min_hamming_bp']}"])
    table(doc, ['Donor / NR4A3 acceptor', 'Class', "Donor bases: ASO sequence (5' to 3')", 'Gap / diff'], data, [2.1,.45,3.65,.8], size=9)

report = {}
for source_name, output_name, supp in [('manuscript.md','ASO-brief-communication.docx',False), ('supporting-material.md','ASO-supporting-material.docx',True)]:
    source = body(ROOT / source_name)
    doc = document(supp)
    small = False
    for block in source.split('\n\n'):
        if block == '{{TABLE1}}': add_table1(doc)
        elif block == '{{EVIDENCE_TABLE}}': add_evidence(doc)
        elif block == '{{CATALOGUE_TABLE}}': add_catalogue(doc)
        elif block.startswith('### '):
            p = para(doc, block[4:], True)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(1)
            for r in p.runs: r.bold = True
        elif block.startswith('# '):
            doc.add_paragraph(block[2:], 'Title')
            doc.core_properties.title = block[2:]
        elif block.startswith('## '):
            title = block[3:]
            small = not supp and title in ('Statements and Declarations','References')
            doc.add_paragraph(title, 'Heading 1')
        else:
            para(doc, block, small or block.startswith(('Independent researcher', 'Keywords:', 'Tristan D. McRae |')))
    doc.save(ROOT / output_name)
    report[output_name] = {'source_sha256':digest(ROOT / source_name), 'docx_sha256':digest(ROOT / output_name), 'tables':len(doc.tables)}
main = body(ROOT / 'manuscript.md')
report['main_words_whitespace'] = len(main.split('## Main text\n\n')[1].split('\n\n{{TABLE1}}')[0].split())
report['abstract_words_whitespace'] = len(main.split('## Abstract\n\n')[1].split('\n\nKeywords:')[0].split())
report['main_reference_count'] = len(re.findall(r'^\[\d+\]', main, re.M))
report['table1_source'] = {'path':str((CAT/'results/all-designs.tsv').relative_to(BASE)), 'sha256':digest(CAT/'results/all-designs.tsv')}
dump(ROOT / 'build-report.json', report)

data_dir = ROOT / 'data'
data_dir.mkdir(exist_ok=True)
names = ['junction-catalogue.tsv', 'all-designs.tsv', 'gap-match-locations.tsv', 'nearest-normal-windows.tsv', 'normal-reference-manifest.tsv', 'normal-reference-transcripts.fasta']
manifest = []
for name in names + ['junction-evidence.tsv']:
    src = (OLD if name == 'junction-evidence.tsv' else CAT) / 'results' / name
    dest = data_dir / name
    shutil.copyfile(src, dest)
    assert digest(src) == digest(dest)
    manifest.append({'file':name, 'bytes':dest.stat().st_size, 'sha256':digest(dest), 'source_commit':SHA, 'source_path':src.relative_to(BASE).as_posix()})
dump(data_dir / 'data-manifest.json', manifest)

# Frozen numerical inputs plus textual source audit. Exclude publishers' full
# articles/images/attachments: those remain linked primary sources in the repo.
files = list(CAT.rglob('*'))
old_files = [p for p in OLD.rglob('*') if p.is_file() and p.suffix.lower() not in {'.pdf','.png','.zip','.xml','.html'}]
files += old_files
files = sorted(p for p in files if p.is_file() and '__pycache__' not in p.parts)
archive_manifest = [{'path':p.relative_to(BASE).as_posix(),'bytes':p.stat().st_size,'sha256':digest(p)} for p in files]
instructions = ('Frozen accepted analysis at commit ' + SHA + '.\n'
    'Extract to an empty directory. Run:\n'
    'python 2026-09-30-full-catalogue/analyze_catalogue.py --output regenerated\n'
    'No network or third-party Python packages are needed for this analysis.\n'
    'Compare regenerated outputs with the accepted results directory.\n'
    'The historical 2026-09-26/evidence/analyze.py has top-level writes; DO NOT run it.\n'
    'Published papers/images are linked, not redistributed in this archive.\n')
with zipfile.ZipFile(ROOT/'reproducibility.zip','w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
    for p in files:
        info = zipfile.ZipInfo(p.relative_to(BASE).as_posix(), date_time=(2026,9,30,0,0,0))
        info.compress_type = zipfile.ZIP_DEFLATED
        z.writestr(info,p.read_bytes())
    for name, text in [('MANIFEST.json',json.dumps(archive_manifest,indent=2)+'\n'), ('REPRODUCE.txt',instructions)]:
        info = zipfile.ZipInfo(name, date_time=(2026,9,30,0,0,0))
        info.compress_type = zipfile.ZIP_DEFLATED
        z.writestr(info,text)
dump(ROOT/'archive-manifest.json',archive_manifest)
print(json.dumps({'main_words':report['main_words_whitespace'], 'abstract_words':report['abstract_words_whitespace'], 'archive_bytes':(ROOT/'reproducibility.zip').stat().st_size, 'data_files':len(manifest)},indent=2))
