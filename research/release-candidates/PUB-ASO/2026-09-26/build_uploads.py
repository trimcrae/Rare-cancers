"""Create journal Word files and a data-derived graphical abstract."""
from pathlib import Path
import csv, json, re, hashlib, zipfile
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'submission'; OUT.mkdir(exist_ok=True)
def plain(s):
    return re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',s).replace('**','').replace('`','')
def runs(p,s):
    for i,part in enumerate(re.split(r'(\*\*.*?\*\*|`[^`]+`)',s)):
        r=p.add_run(part.strip('*`'))
        if part.startswith('**'):r.bold=True
        if part.startswith('`'):r.font.name='Consolas';r.font.size=Pt(10)
def field(p,name):
    f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),name);p._p.append(f)
def build(md,filename,double=True):
    if md.startswith('---\n'): md=md.split('---\n',2)[2].lstrip()
    # CBC's recorded official guide places manuscript tables and captions at the end.
    if double:
        lines=md.splitlines(); body=[]; tables=[]; i=0
        while i<len(lines):
            if lines[i].startswith('**Table '):
                block=[lines[i]];i+=1
                while i<len(lines) and not lines[i].strip():i+=1
                while i<len(lines) and lines[i].strip().startswith('|'):
                    block.append(lines[i]);i+=1
                tables.append('\n'.join(block))
            else:body.append(lines[i]);i+=1
        md='\n'.join(body)+'\n\n'+'\n\n'.join('<!-- table-page -->\n'+t for t in tables)
    doc=Document();sec=doc.sections[0]
    doc.settings.odd_and_even_pages_header_footer=False
    sec.different_first_page_header_footer=False
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Cm(3)
    sec.left_margin=sec.right_margin=Cm(3)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name];style.font.name='Times New Roman'
        style.font.color.rgb=RGBColor(0,0,0)
        fonts=style.element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is not None:
            for key in list(fonts.attrib):
                if key.endswith('Theme'):del fonts.attrib[key]
            for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    normal=doc.styles['Normal'];normal.font.size=Pt(12 if double else 11)
    normal.paragraph_format.line_spacing=2 if double else 1.15
    normal.paragraph_format.space_after=Pt(5 if double else 6)
    doc.styles['Title'].font.size=Pt(18)
    doc.styles['Title'].paragraph_format.line_spacing=1.1
    for name,size in [('Heading 1',13),('Heading 2',12)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.line_spacing=1.1
        doc.styles[name].paragraph_format.space_before=Pt(12)
        doc.styles[name].paragraph_format.space_after=Pt(6)
    head=sec.header.paragraphs[0];head.text='EMC junction provenance'
    head.style='Normal';head.runs[0].font.size=Pt(9);head.paragraph_format.line_spacing=1
    foot=sec.footer.paragraphs[0];foot.alignment=2;foot.paragraph_format.line_spacing=1
    foot.add_run('Page ');field(foot,'PAGE')
    lines=md.splitlines();i=0;code=False;front=True;refs=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('```'):code=not code;continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
            cells=[[plain(s.strip()) for s in r.strip('|').split('|')] for r in rows]
            cells=[r for r in cells if not all(re.fullmatch('[: -]+',s) for s in r)]
            table=doc.add_table(rows=0,cols=len(cells[0]));table.autofit=False
            usable=8.5-6/2.54
            if len(cells[0])==4: widths=[1.65,1.35,1.55,usable-4.55]
            else: widths=[usable/len(cells[0])]*len(cells[0])
            for c,w in zip(table.columns,widths):c.width=Inches(w)
            pr=table._tbl.tblPr
            borders=OxmlElement('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single' if side in ['top','bottom','insideH'] else 'nil');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders)
            for ri,row in enumerate(cells):
                rr=table.add_row()
                trpr=rr._tr.get_or_add_trPr()
                nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
                if ri==0:trpr.append(OxmlElement('w:tblHeader'))
                for ci,(cell,text) in enumerate(zip(rr.cells,row)):
                    cell.width=Inches(widths[ci]);cell.vertical_alignment=1
                    cp=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
                    for side in ['top','bottom','left','right']:
                        el=OxmlElement('w:'+side);el.set(qn('w:w'),'70');el.set(qn('w:type'),'dxa');mar.append(el)
                    cp.append(mar)
                    p=cell.paragraphs[0];p.paragraph_format.line_spacing=1.08;p.paragraph_format.space_after=Pt(3)
                    r=p.add_run(text);r.font.size=Pt(10);r.bold=ri==0
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            continue
        if line=='<!-- table-page -->':
            doc.add_page_break();continue
        if line.startswith('# '):p=doc.add_paragraph(line[2:],'Title')
        elif line.startswith('### '):p=doc.add_paragraph(line[4:],'Heading 2')
        elif line.startswith('## '):
            heading=line[3:];p=doc.add_paragraph(heading,'Heading 1')
            if double and heading=='Abstract':p.paragraph_format.page_break_before=True
            if heading=='1. Introduction':
                front=False
                if double:p.paragraph_format.page_break_before=True
            if heading=='References':refs=True
        else:
            p=doc.add_paragraph()
            runs(p,plain(line) if '[' in line and '](' in line else line)
            if code:
                p.paragraph_format.line_spacing=1
                for r in p.runs:r.font.name='Consolas';r.font.size=Pt(10)
            if line.startswith('**Table'):p.paragraph_format.keep_with_next=True
        if double and not line.startswith('#'):p.paragraph_format.line_spacing=2
    doc.core_properties.author='Tristan D. McRae'
    doc.core_properties.title=md.splitlines()[0].lstrip('# ')
    doc.save(OUT/filename)
build((ROOT/'manuscript.md').read_text(encoding='utf-8'),'ASO-manuscript.docx')
build((ROOT/'supplementary-methods.md').read_text(encoding='utf-8'),'ASO-supplementary-methods.docx',False)

rows=list(csv.DictReader((ROOT/'evidence/results/design-comparisons.tsv').open(),delimiter='\t'))
tc=sorted([r for r in rows if r['junction']=='TCF12_e5__NR4A3_e3'],key=lambda r:int(r['donor_bases']))
fig=plt.figure(figsize=(13.28,5.31),dpi=100,facecolor='white')
fig.text(.045,.92,'Transcript provenance changes EMC antisense comparisons',fontsize=20,weight='bold')
fig.text(.045,.79,'USZ20 model annotation',fontsize=15,weight='bold')
fig.text(.045,.69,'EWSR1 e13 / NR4A3 e2\nENST00000414183.2 / ENST00000330847.1',fontsize=12,linespacing=1.5)
fig.text(.045,.55,'Same reference boundaries',fontsize=13,color='#1B607A')
fig.text(.045,.44,'EWSR1 e12 / NR4A3 e3\nNM_005243.4 / NM_006981.4',fontsize=12,linespacing=1.5)
fig.text(.045,.27,'USZ20: reference reconstruction\nUSZ22: exact RNA junction unresolved',fontsize=12,linespacing=1.5)
ax=fig.add_axes([.57,.25,.39,.5])
x=[int(r['donor_bases']) for r in tc]
ax.plot(x,[int(r['legacy_six_parent_longest_gap_spanning_match_bp']) for r in tc],'o--',color='#6B6B6B',label='Original 6 transcripts',linewidth=2)
ax.plot(x,[int(r['expanded_parent_corpus_longest_match_bp']) for r in tc],'o-',color='#176482',label='Expanded 77 records',linewidth=2)
ax.set_title('TCF12 junction: five binding positions',fontsize=14,loc='left',pad=12)
ax.set_xlabel('Donor bases in the 16-base target',fontsize=11)
ax.set_ylabel('Longest gap-spanning match (nt)',fontsize=11)
ax.set_xticks(x);ax.set_ylim(5.8,14);ax.set_yticks([6,8,10,12,14])
ax.spines[['top','right']].set_visible(False);ax.grid(axis='y',alpha=.18);ax.legend(frameon=False,fontsize=10)
fig.text(.045,.075,'9 of 35 source-linked designs gain longer matches with added parent isoforms.',fontsize=14,weight='bold')
fig.text(.045,.025,'Sequence comparisons do not establish RNA cleavage or biological selectivity.',fontsize=11)
fig.savefig(OUT/'ASO-graphical-abstract.png',dpi=100)
fig.savefig(OUT/'ASO-graphical-abstract.pdf')
plt.close(fig)
for name in ['highlights.txt','graphical-abstract-caption.txt']:
    (OUT/name).write_bytes((ROOT/name).read_bytes())
# Compact journal dataset. Exact original large read prefixes remain in local evidence.
with zipfile.ZipFile(OUT/'ASO-supplementary-data.zip','w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
    for p in sorted((ROOT/'evidence').rglob('*')):
        if p.is_file() and '__pycache__' not in p.parts:z.write(p,'evidence/'+p.relative_to(ROOT/'evidence').as_posix())
    z.write(ROOT/'evidence-manifest.json','evidence-manifest.json')
manifest=[dict(file=p.name,bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in sorted(OUT.iterdir()) if p.is_file() and p.name!='build-manifest.json']
(OUT/'build-manifest.json').write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'outputs':len(manifest),'dataset_bytes':(OUT/'ASO-supplementary-data.zip').stat().st_size}))
