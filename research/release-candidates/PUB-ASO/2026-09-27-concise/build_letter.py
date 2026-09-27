"""Build the concise author-review letter; no scientific computation or networking."""
from pathlib import Path
import hashlib
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
def manuscript_body(source):
    if source.startswith('---\n'):
        _, separator, body = source[4:].partition('\n---\n')
        if not separator:
            raise ValueError('Unclosed repository frontmatter in letter.md')
        return body.lstrip('\n')
    return source


source_text = (ROOT / 'letter.md').read_text(encoding='utf-8')
text = manuscript_body(source_text)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.75)
sec.left_margin = sec.right_margin = Inches(.85)
for name in ['Normal', 'Title', 'Heading 1']:
    s = doc.styles[name]
    s.font.name = 'Times New Roman'
    s.font.color.rgb = RGBColor(0, 0, 0)
    for e in list(s.element.iter(qn('w:pBdr'))):
        e.getparent().remove(e)
    rpr = s.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is not None:
        for k in list(fonts.attrib):
            if k.endswith('Theme'):
                del fonts.attrib[k]
        for k in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:' + k), 'Times New Roman')
normal = doc.styles['Normal']
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
title = doc.styles['Title']
title.font.size = Pt(17)
title.paragraph_format.line_spacing = 1
title.paragraph_format.space_after = Pt(7)
heading = doc.styles['Heading 1']
heading.font.size = Pt(11)
heading.font.bold = True
heading.paragraph_format.line_spacing = 1
heading.paragraph_format.space_before = Pt(6)
heading.paragraph_format.space_after = Pt(3)
heading.paragraph_format.keep_with_next = True
small = False
for block in text.strip().split('\n\n'):
    if block.startswith('# '):
        p = doc.add_paragraph(block[2:], 'Title')
    elif block.startswith('## '):
        p = doc.add_paragraph(block[3:], 'Heading 1')
        small = True
    else:
        p = doc.add_paragraph(block)
        if small or block.startswith('Independent researcher'):
            p.paragraph_format.line_spacing = 1.05
            for r in p.runs:
                r.font.size = Pt(10.5)
        if block == 'Tristan D. McRae':
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
        if block == 'To the Editor,':
            p.paragraph_format.keep_with_next = True
doc.core_properties.title = text.splitlines()[0][2:]
doc.core_properties.author = 'Tristan D. McRae'
doc.core_properties.subject = 'EMC fusion junction sequence provenance and normal-parent comparison'
out = ROOT / 'ASO-letter.docx'
doc.save(out)
body = text.split('To the Editor,\n\n', 1)[1].split('\n\n## Declarations', 1)[0]
report = {
    'main_words_whitespace': len(body.split()),
    'complete_document_words_whitespace': sum(len(p.text.split()) for p in doc.paragraphs),
    'references': len(re.findall(r'^\[\d+\]', text, re.M)),
    'tables': len(doc.tables),
    'figures': len(doc.inline_shapes),
    'source_sha256': hashlib.sha256((ROOT / 'letter.md').read_bytes()).hexdigest(),
    'manuscript_body_sha256': hashlib.sha256(text.encode()).hexdigest(),
    'docx_sha256': hashlib.sha256(out.read_bytes()).hexdigest(),
    'scientific_inputs_changed': False,
    'page_count_and_visual_verification': 'require canonical render and all-page inspection',
    'layout': 'US Letter; 0.75 inch vertical and 0.85 inch horizontal margins; 12 point Times New Roman, 1.15 line spacing; 10.5 point declarations and references',
}
(ROOT / 'build-report.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8')
print(json.dumps(report, indent=2))
