"""Build the Cancer Genetics editable package from committed Markdown; no science runs."""
from pathlib import Path
import re,json,hashlib,sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[3]
HERE=Path(__file__).resolve().parent
STEM='emc-tissue-rna-prioritization'
BLACK=RGBColor(0,0,0)
def text_part(p,s):
    s=re.sub(r'\[([^]]+)\]\((https?://[^)]+)\)',r'\1 (\2)',s)
    for part in re.split(r'(\*\*.*?\*\*|\*[^*]+\*|\x60[^\x60]+\x60|https?://[^\s]+)',s):
        if not part:continue
        if part.startswith('http'):
            trail=''
            while part and part[-1] in '.,;)':trail=part[-1]+trail;part=part[:-1]
            rid=p.part.relate_to(part,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid)
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'000000');pr.append(c)
            size=OxmlElement('w:sz');size.set(qn('w:val'),'18');pr.append(size);r.append(pr)
            t=OxmlElement('w:t');t.text=part;r.append(t);h.append(r);p._p.append(h)
            if trail:p.add_run(trail)
        else:
            r=p.add_run(part.strip('*').strip(chr(96)))
            if part.startswith('**'):r.bold=True
            elif part.startswith('*'):r.italic=True
            elif part.startswith(chr(96)):r.font.name='Consolas';r.font.size=Pt(9)
def base_doc(kind):
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.27);s.page_height=Inches(11.69)
    s.top_margin=s.bottom_margin=Inches(.8);s.left_margin=s.right_margin=Inches(.8)
    s.header_distance=s.footer_distance=Inches(.35)
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12 if kind=='main' else 11)
    normal.font.color.rgb=BLACK;normal.paragraph_format.line_spacing=1.5 if kind=='main' else 1.15
    normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',16),('Heading 1',13),('Heading 2',12),('Heading 3',11)]:
        st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=BLACK;st.font.bold=True
        st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
    for st in d.styles:
        for el in list(st.element.xpath('.//w:pBdr')):el.getparent().remove(el)
        for el in st.element.xpath('.//w:rFonts'):
            for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
                el.attrib.pop(qn('w:'+attr),None)
    d.core_properties.author='Tristan D. McRae';d.core_properties.subject='Computational tissue RNA study prepared for Cancer Genetics'
    f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=f.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
    return d

def table(d,lines):
    rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',c or '-') for c in r)]
    n=len(rows[0]);assert all(len(r)==n for r in rows)
    t=d.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths={2:[1.5,5.17],3:[1.25,1.5,3.92],4:[.85,2.03,1.7,2.09],6:[.75,1.14,1.14,1.06,1.32,1.26]}.get(n,[6.67/n]*n)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,(cell,val,w) in enumerate(zip(cells,row,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.05
            text_part(p,val)
            for run in p.runs:run.font.size=Pt(9.5);run.bold=(i==0)
            if j>0 and n==6:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');pr.append(sh)
        trpr=t.rows[-1]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
        if i==0:rep=OxmlElement('w:tblHeader');trpr.append(rep)
    d.add_paragraph()

def from_md(md,kind):
    d=base_doc(kind);md=re.sub(r'^---\n.*?\n---\n','',md,flags=re.S).strip();lines=md.splitlines();i=0;section='';in_code=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('~~~') or line.startswith(chr(96)*3):in_code=not in_code;continue
        if line.startswith('|'):
            chunk=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):chunk.append(lines[i].strip());i+=1
            table(d,chunk);continue
        if line.startswith('# '):
            title=line[2:];d.add_paragraph(title,'Title');d.core_properties.title=title;continue
        if line.startswith('## '):
            section=line[3:]
            if kind=='main' and section=='Introduction':d.add_page_break()
            d.add_paragraph(section,'Heading 1');continue
        if line.startswith('### '):d.add_paragraph(line[4:],'Heading 2');continue
        p=d.add_paragraph();text_part(p,line)
        if in_code:
            for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
        if section in ['Abstract','Keywords','References'] or (kind=='main' and not section):
            p.paragraph_format.line_spacing=1.1
            for r in p.runs:r.font.size=Pt(10.5)
        if line.startswith('[') and section=='References':
            p.paragraph_format.space_after=Pt(8);p.paragraph_format.keep_together=True
    return d

main=(HERE/(STEM+'.md')).read_text(encoding='utf8');si=(HERE/(STEM+'-si.md')).read_text(encoding='utf8')
d=from_md(main,'main')
for i in [1,2]:
    d.add_page_break();d.add_paragraph('Figure '+str(i),'Heading 1')
    p=d.add_paragraph();p.paragraph_format.line_spacing=1;p.add_run().add_picture(str(ROOT/f'research/manuscripts/figures/surface-tissue-rna-figure{i}.png'),width=Inches(6.67))
d.save(HERE/(STEM+'-manuscript.docx'))
s=from_md(si,'si');s.save(HERE/(STEM+'-supplementary-information.docx'))
cover='''# Cancer Genetics cover letter

Dear Editors,

Please consider “CSPG4 tissue RNA enrichment in extraskeletal myxoid chondrosarcoma depends on comparator and sequencing year” as an Original Research Article in Cancer Genetics.

This secondary analysis characterizes a fixed 11-gene panel using publicly released tissue RNA measurements. Nine primary EMC specimens provide marginal contrasts; only four unique EMC patients contribute to the three sequencing-year-matched comparisons. Original GSE24369 arrays provide a separate same-histology comparison. CSPG4 alone meets the frozen allocation rule, while removing sequencing-year 2019 reverses the matched composite and comparison with dermatofibrosarcoma protuberans (DFSP) is nonpositive. The contribution is a reproducible account of comparator-dependent molecular evidence and a specific question for tissue localization, without a claim of validated protein expression or clinical utility.

The manuscript is relevant to comparative molecular characterization of cancer. It reports all panel results and retains sparse-stratum uncertainty, incomplete historical independence and normal-expression limitations. The full source/code archive, frozen protocols and effect tables are publicly available with Research Square version 1 at https://doi.org/10.21203/rs.3.rs-10959636/v1. The earlier aiXiv posting is https://aixiv.science/abs/aixiv.260907.000003. This journal-tailored revision retains the analyses and scientific figures and improves presentation, references and data links.

The research received no funding, and the author declares no competing interests. It reanalyzes public data and reports no new participant recruitment, specimens, intervention or laboratory experiment; it asserts no new ethics approval, exemption or waiver. OpenAI Codex assistance with code, computation, verification and drafting is disclosed in the manuscript. The author is responsible for its final interpretation.

The requested publication route is standard subscription, with online-only color and no optional open-access or accelerated service.

Sincerely,

Tristan D. McRae
Independent researcher, unaffiliated
trimcrae@gmail.com
ORCID 0000-0002-1823-1451
'''
(HERE/(STEM+'-cover-letter.md')).write_text("---\nid: DOC-SURFACE-CANCER-GENETICS-COVER-LETTER\ntitle: Cancer Genetics cover letter\nkind: letter\nstatus: live\npurpose: Prepare an accurate cover letter for the selected journal.\nscope: Journal preparation; no journal submission or acceptance asserted.\naudience: [author, journal editors]\ndate: 2026-09-09\nlast_verified: 2026-09-09\n---\n\n"+cover,encoding='utf8',newline='\n')
c=from_md(cover,'cover');c.save(HERE/(STEM+'-cover-letter.docx'))
highlights=['CSPG4 alone met the frozen prioritization rule within an 11-gene RNA panel.','The same LGFMS comparison supported CSPG4 ordering in two tissue cohorts.','Only four unique EMC patients contributed to the matched comparisons.','Removing sequencing-year 2019 reversed the matched CSPG4 composite.','The data motivate protein localization studies rather than clinical use.']
assert 3<=len(highlights)<=5 and all(len(x)<=85 for x in highlights)
(HERE/(STEM+'-highlights.txt')).write_text('\n'.join(highlights)+'\n',encoding='utf8')
(HERE/(STEM+'-figure-legends.txt')).write_text(main.split('## Figure legends\n',1)[1].strip()+'\n',encoding='utf8')
outputs=[]
for suffix in ['-manuscript.docx','-supplementary-information.docx','-cover-letter.docx','-cover-letter.md','-highlights.txt','-figure-legends.txt']:
    p=HERE/(STEM+suffix);outputs.append({'path':str(p.relative_to(ROOT)).replace('\\','/'),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
print(json.dumps({'outputs':outputs,'highlights_characters':[len(x) for x in highlights],'editable_si_tables':len(s.tables),'main_figures':len(d.inline_shapes),'science_runs':0}))
