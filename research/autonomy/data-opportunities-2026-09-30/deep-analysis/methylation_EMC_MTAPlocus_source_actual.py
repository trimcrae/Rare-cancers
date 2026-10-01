import csv,gzip,hashlib,io,json,pathlib,re,urllib.parse,urllib.request,xml.etree.ElementTree as ET,zipfile
import openpyxl
from pypdf import PdfReader
OUT=pathlib.Path('campaign-output/methylation-emc-MTAP');OUT.mkdir(parents=True,exist_ok=True)
MD_URL='https://raw.githubusercontent.com/qtran1/MeQTrack_app/578bb615e0e89858ffb02cbe1da87680c74a01ba/reference/GSE140686_sarcoma_methylation_labels.csv';MD_SHA='c3b102c078a1035f4589952ea8b50ff3a8fbbb63ed8edd425aed2ea343119a19';BASE='https://ftp.ncbi.nlm.nih.gov/geo/series/GSE140nnn/GSE140686/';receipts=[]
def sha(b):return hashlib.sha256(b).hexdigest()
def get(url,limit=10000000):
 url=url.replace('ftp://ftp.ncbi.nlm.nih.gov/','https://ftp.ncbi.nlm.nih.gov/');req=urllib.request.Request(url,headers={'User-Agent':'Rare-cancers-finite-measured-source-audit/1'})
 with urllib.request.urlopen(req,timeout=120) as r:b=r.read(limit+1);status=r.status;final=r.url
 if len(b)>limit:raise ValueError('Declaredsourceboundexceeded')
 receipts.append({'url':url,'final_url':final,'http_status':status,'bytes':len(b),'sha256':sha(b)});return b
md=get(MD_URL,500000);assert sha(md)==MD_SHA;rows=list(csv.DictReader(io.StringIO(md.decode())));queries=[r for r in rows if r['Meth_Class']=='EMCS'];controls=[r for r in rows if r['Meth_Class'] in ['CTRL-MUS','CTRL-REA'] and r['DNA']=='FFPE'];assert len(queries)==10 and len(controls)==18;frozen=[dict(r,analysis_role='EMC_query' if r in queries else 'FFPE_non_neoplastic_reference') for r in queries+controls];xml=get('https://www.ebi.ac.uk/europepmc/webservices/rest/PMC7819999/fullTextXML',1000000);(OUT/'primary.xml').write_bytes(xml);tree=ET.fromstring(xml);passages=[' '.join(' '.join(p.itertext()).split()) for p in tree.findall('.//p') if re.search(r'copy.number|conumee|deletion|MTAP|CDKN2',' '.join(p.itertext()),re.I)];supp=[];seen=set();static='https://static-content.springer.com/esm/art%3A10.1038%2Fs41467-020-20603-4/MediaObjects/'
for n in tree.findall('.//supplementary-material'):
 caption=' '.join(' '.join(n.itertext()).split())
 for a in n.iter():
  for k,v in a.attrib.items():
   if not k.endswith('href'):continue
   url=v if v.startswith('http') else static+urllib.parse.quote(v)
   if url in seen:continue
   seen.add(url);rec={'url':url,'caption':caption}
   try:
    b=get(url,60000000);name=pathlib.PurePosixPath(urllib.parse.urlparse(url).path).name;(OUT/name).write_bytes(b);rec.update(bytes=len(b),sha256=sha(b))
    if name.lower().endswith('.xlsx'):
     wb=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True);rec['sheets']=[]
     for ws in wb.worksheets:
      rr=list(ws.iter_rows(values_only=True));hits=[]
      for i,row in enumerate(rr,1):
       literal=[None if v is None else v if isinstance(v,(str,int,float,bool)) else str(v) for v in row];st=' '.join(str(v or '') for v in literal)
       if re.search(r'\bMTAP\b|\bCDKN2[AB]\b|copy.number|segment|CNV|deletion',st,re.I):hits.append({'excel_row':i,'cells':literal})
      rec['sheets'].append({'name':ws.title,'rows':len(rr),'columns':ws.max_column,'first_12_rows':[[None if v is None else str(v) for v in r] for r in rr[:12]],'all_locus_and_copy_number_keyword_rows':hits})
    elif name.lower().endswith('.pdf'):
     pages=[p.extract_text() or '' for p in PdfReader(io.BytesIO(b)).pages];(OUT/(name+'.txt')).write_text('\n\n'.join(pages));rec['pages']=len(pages);rec['copy_number_or_locus_pages']=[{'page':i+1,'text':s} for i,s in enumerate(pages) if re.search(r'\bMTAP\b|\bCDKN2[AB]\b|copy.number|conumee|deletion',s,re.I)]
    else:rec['status']='Preserved; format-specific followthrough ifrelevant'
   except Exception as e:rec['error']=type(e).__name__+': '+str(e)
   supp.append(rec)
geo={}
for platform in ['GPL13534','GPL21145']:
 b=get(BASE+'matrix/GSE140686-'+platform+'_series_matrix.txt.gz',2000000);s=gzip.decompress(b).decode();fields={}
 for line in s.splitlines():
  if not line.startswith('!Sample_'):continue
  cells=next(csv.reader([line],delimiter='\t'));fields.setdefault(cells[0],[]).append(cells[1:])
 gsms=fields['!Sample_geo_accession'][0];titles=fields['!Sample_title'][0];links=fields.get('!Sample_supplementary_file',[])
 for i,(gsm,title) in enumerate(zip(gsms,titles)):
  z=re.search(r'reference case\s+(\d+)',title,re.I)
  if z:geo['REFERENCE_SAMPLE '+str(int(z.group(1)))]={'gsm':gsm,'title':title,'platform':platform,'supplementary_urls':[l[i] for l in links if i<len(l) and l[i] not in ['NONE','']]}
raw=[];idatdir=OUT/'idats';idatdir.mkdir(exist_ok=True)
for r in frozen:
 rec=dict(r);g=geo.get(r['ID']);rec['geo']=g
 if not g:rec['error']='Exactreference-GSMmappingabsent';raw.append(rec);continue
 urls=[u for u in g['supplementary_urls'] if re.search(r'\.idat(?:\.gz)?$',u,re.I)]
 if not urls:
  u='https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc='+g['gsm']+'&targ=self&form=text&view=full'
  try:
   s=get(u,3000000).decode();(OUT/(g['gsm']+'.soft.txt')).write_text(s);urls=[l.split('=',1)[1].strip() for l in s.splitlines() if l.startswith('!Sample_supplementary_file =') and re.search(r'\.idat(?:\.gz)?$',l,re.I)]
  except Exception as e:rec['sample_metadata_error']=type(e).__name__+': '+str(e)
 rec['channels']=[]
 for channel in ['Grn','Red']:
  hits=[u for u in urls if r['IDAT'] in u and re.search('_'+channel+r'\.idat(?:\.gz)?$',u,re.I)];ch={'channel':channel,'matching_declared_urls':hits}
  if len(hits)!=1:ch['error']='ExactdeclaredIDAT/channelURLnotunique'
  else:
   try:
    b=get(hits[0],25000000);unpacked=gzip.decompress(b) if b.startswith(b'\x1f\x8b') else b;assert unpacked[:4]==b'IDAT' and len(unpacked)<=60000000;filename=r['IDAT']+'_'+channel+'.idat.gz';packed=b if b.startswith(b'\x1f\x8b') else gzip.compress(b,mtime=0);(idatdir/filename).write_bytes(packed);ch.update(file=filename,compressed_bytes=len(packed),compressed_sha256=sha(packed),uncompressed_bytes=len(unpacked),uncompressed_sha256=sha(unpacked),status='acquired')
   except Exception as e:ch['error']=type(e).__name__+': '+str(e)
  rec['channels'].append(ch)
 raw.append(rec);print('EMC_MTAP_RAW_PROFILE '+r['ID'],flush=True)
archive={'url':BASE+'suppl/GSE140686_RAW.tar'}
if any(any(ch.get('status')!='acquired' for ch in r.get('channels',[])) or len(r.get('channels',[]))!=2 for r in raw):
 try:
  req=urllib.request.Request(archive['url'],headers={'Range':'bytes=0-511','User-Agent':'Rare-cancers-finite-measured-source-audit/1'})
  with urllib.request.urlopen(req,timeout=120) as r:b=r.read(513);archive.update(http_status=r.status,content_range=r.headers.get('Content-Range'),bytes_read=len(b))
  assert archive['http_status']==206 and len(b)==512;archive['first_member_name']=b[:100].split(b'\0',1)[0].decode(errors='replace');archive['header_sha256']=sha(b);archive['status']='Range-readable; exactselectedoffsetrecovery runnable'
 except Exception as e:archive['error']=type(e).__name__+': '+str(e)
complete=[r for r in raw if len(r.get('channels',[]))==2 and all(c.get('status')=='acquired' for c in r['channels'])]
obj={'schema':'emc-ten-reference-MTAP-primary-and-raw-source/1','frozen_before_CN_values':True,'genes':['MTAP','CDKN2A','CDKN2B'],'metadata_sha256':MD_SHA,'primary_xml_sha256':sha(xml),'primary_copy_number_passages':passages,'primary_supplement_inventory':supp,'query_n':10,'reference_n':18,'raw_profiles':raw,'complete_pairs':len(complete),'complete_query_pairs':sum(r['analysis_role']=='EMC_query' for r in complete),'complete_reference_pairs':sum(r['analysis_role']!='EMC_query' for r in complete),'raw_archive':archive,'source_receipts':receipts,'analysis_status':'Sourceacquisition; literalCNcallsadjudication/conumee stillpending','interpretation':['Betavalues notCN.','ReferenceCNneutrality unverified; retainLOOcontrols.','9target450K/1EPIC; commonprobesprimary.','Relativearrayratios notabsolute/biallelicloss/MTAPprotein/MTA/dependency/drugresponse.']}
(OUT/'source-and-IDAT-manifest.json').write_text(json.dumps(obj,indent=2,ensure_ascii=False,allow_nan=False)+'\n')
with zipfile.ZipFile(OUT/'targeted-EMC-and-FFPE-reference-IDATs.zip','w',compression=zipfile.ZIP_STORED) as z:
 for p in sorted(idatdir.glob('*.idat.gz')):z.write(p,'idats/'+p.name)
 z.write(OUT/'source-and-IDAT-manifest.json','source-and-IDAT-manifest.json')
print('EMC_MTAP_RAW_SOURCE_BEGIN');print(json.dumps(obj,separators=(',',':'),ensure_ascii=False,allow_nan=False));print('EMC_MTAP_RAW_SOURCE_END');raise SystemExit(0 if len(complete)==28 else 1)
