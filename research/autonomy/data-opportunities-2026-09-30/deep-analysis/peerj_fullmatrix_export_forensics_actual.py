import json,hashlib,re,datetime,collections
from pathlib import Path
import numpy as np
from scipy.stats import rankdata
from openpyxl import load_workbook
import peerj_expression_diagnostic_actual as p
import spatial_marker_followthrough_actual as sm
import scientific_source_access_followthrough as src
OUT=sm.OUT;SEED=2026100139;SETTINGS=[('unadjusted',[]),('repeated_value_burden',['repeated_value_fraction']),('global_expression_median',['published_log2CPM_median']),('both_endogenous_export_measures',['repeated_value_fraction','published_log2CPM_median'])]
def corr(x,y):
 x=np.asarray(x,float)-np.mean(x);y=np.asarray(y,float)-np.mean(y);den=np.linalg.norm(x)*np.linalg.norm(y)
 return float(np.dot(x,y)/den) if den>1e-12 else None
def model(x,y,z):
 n=len(x);xr=rankdata(x,method='average');yr=rankdata(y,method='average');C=np.column_stack([np.ones(n)]+[rankdata(z[:,j],method='average') for j in range(z.shape[1])]);u,s,v=np.linalg.svd(C,full_matrices=False);rank=int(np.sum(s>max(C.shape)*s[0]*np.finfo(float).eps));Q=u[:,:rank];M=np.eye(n)-Q@Q.T;xx=M@xr;yy=M@yr
 return corr(xx,yy),M,xx,yy,rank
def test(x,y,z,labels,setting,left,right,seed):
 rho,M,xx,yy,rank=model(x,y,z);n=len(x);r={'setting':setting,'left':left,'right':right,'n_source_specimens':n,'nuisance_design_rank':rank,'residual_dimensions':n-rank,'partial_rank_correlation':rho,'nuisance_labels':labels,'no_linked_outcomes':True}
 if rho is None or n-rank<4:r['status']='withheldconstant_or_insufficient_residual_dimensions';return r
 rng=np.random.default_rng(seed);perms=np.array([rng.permutation(n) for unused in range(10000)]);null_y=yy[perms]@M.T;den=np.linalg.norm(xx)*np.linalg.norm(null_y,axis=1);good=den>1e-12;null=(null_y[good]@xx)/den[good];assert len(null)==10000,'Constantpermutedresidual draws require explicitamendment';r.update(status='measured',approximate_conditional_residual_permutation_p=float((1+np.sum(np.abs(null)>=abs(rho)-1e-14))/(1+len(null))),permutation_finite_draws=len(null),permutation_null_central95=np.quantile(null,[.025,.975]).tolist(),permutation_interpretation='Freedman-Lane style fixedmidrank residualperm; approximateconditionalexchangeability, notcausal technical correction')
 boot=[];discarded=0
 for unused in range(2000):
  ix=rng.integers(0,n,n);v,mm,xxx,yyy,rr=model(x[ix],y[ix],z[ix])
  if v is None or n-rr<4:discarded+=1
  else:boot.append(v)
 r.update(bootstrap_requested_draws=2000,bootstrap_finite_draws=len(boot),bootstrap_discarded_draws=discarded,reranked_source_specimen_bootstrap95=np.quantile(boot,[.025,.975]).tolist() if len(boot)>=1800 else None,leave_one_source_specimen_out=[])
 for j,label in enumerate(p.IDS):
  keep=np.arange(n)!=j;v,mm,xxx,yyy,rr=model(x[keep],y[keep],z[keep]);r['leave_one_source_specimen_out'].append({'excluded_specimen':label,'partial_rank_correlation':v,'nuisance_design_rank':rr})
 return r
def bh(rr,key,out,total):
 order=sorted([(i,r[key]) for i,r in enumerate(rr) if r.get(key) is not None],key=lambda x:x[1]);q=1.
 for k in range(len(order)-1,-1,-1):i,pv=order[k];q=min(q,pv*total/(k+1));rr[i][out]=q
def main():
 OUT.mkdir(parents=True,exist_ok=True);path,receipt=p.get();book=src.xlsx_tables(path.read_bytes());sheet=next(t for t in book['tables'] if t['sheet']=='EMC_Gene-expression_Log2CPM');rows=sheet['rows'];assert rows[0]==['symbol']+p.IDS and len(rows)==9501 and all(len(v)==13 for v in rows[1:]);symbols=[str(row[0]).strip() for row in rows[1:]];V=np.array([[float(v) for v in row[1:]] for row in rows[1:]],float);assert V.shape==(9500,12) and np.isfinite(V).all();lookup=collections.defaultdict(list)
 for i,s in enumerate(symbols):lookup[s].append(i)
 selected={};mapping={}
 for gene in p.GENES:
  hits=[(s,i) for s in [gene]+(['H1-10'] if gene=='H1FX' else[]) for i in lookup.get(s,[])];mapping[gene]=hits;assert len(hits)==1;selected[gene]=hits[0][1]
 repeated=np.abs(V-p.FLOOR)<=5e-10;floor=repeated.mean(axis=0);median=np.median(V,axis=0);metrics={'repeated_value_fraction':floor,'published_log2CPM_median':median};numeric=[];wb=load_workbook(path,read_only=True,data_only=True);ws=wb['EMC_Gene-expression_Log2CPM']
 for i,s in enumerate(symbols):
  if not re.fullmatch(r'\d+',s):continue
  cell=ws.cell(i+2,1);serial=int(s);compatible=20000<=serial<=60000;date=(datetime.datetime(1899,12,30)+datetime.timedelta(days=serial)).date().isoformat() if compatible else None;numeric.append({'row_1based':i+2,'literal_symbol':s,'literal_cell_value':str(cell.value),'source_cell_data_type':cell.data_type,'source_number_format':cell.number_format,'Excel1900_date_candidate':date,'source_gene_identity_changed':False,'date_compatibility_is_conversion_proof':False,'twelve_literal_log2CPM_values':V[i].tolist()})
 wb.close();numeric_indices=[i for i,s in enumerate(symbols) if re.fullmatch(r'\d+',s)];mask=np.ones(len(V),bool);mask[numeric_indices]=False;drop_floor=repeated[mask].mean(axis=0);drop_median=np.median(V[mask],axis=0);centered=V-V.mean(axis=1,keepdims=True);sd=centered.std(axis=1);usable=sd>1e-12;standardized=centered[usable]/sd[usable,None];U,S,VT=np.linalg.svd(standardized,full_matrices=False);pc=VT.T*S;fraction=S*S/np.sum(S*S);tests=[]
 for si,(setting,labels) in enumerate(SETTINGS):
  Z=np.column_stack([metrics[label] for label in labels]) if labels else np.empty((12,0))
  for pi,(left,right) in enumerate(p.PAIRS):tests.append(test(V[selected[left]],V[selected[right]],Z,labels,setting,left,right,SEED+si*100+pi))
 bh(tests,'approximate_conditional_residual_permutation_p','BH_q_fixed12_descriptive',12);result={'schema':'emc-PeerJ-fullmatrix-export-forensics/1','executed_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'source_receipt':receipt,'amendment':{'dated_utc':'2026-10-01','selection':'Exactfrozen9500genes×12specimenexpression/all3publishedpostselectiongenepairs','settings':SETTINGS,'nuisance':'Repeated.010026459±5e-10valuefraction/globalpublishedlog2CPMmedian; endogenous export measures, neither verifieddepth/batch norbiologicalconfounder','tests':'All12partialranktests/10000conditionalresidualperms/BH12/2000rerankedresampledraws/all12leaveonespecimenout','identity':'Allliteral numeric/datecompatible rows retained; noMARCHorothergenenames invented','limits':'No outcomejoin, prognosisfit, newselectedsignature, causalbatchcorrection, wetlabprotein evidence; nuisance associations hypothesisgenerating'},'matrix_shape':list(V.shape),'frozen_gene_mapping':mapping,'all_numeric_literal_gene_rows':numeric,'nonconstant_source_genes':int(usable.sum()),'constant_source_genes':int((~usable).sum()),'all12_specimen_export_metrics':[{'specimen':sid,'repeated_value_rows':int(repeated[:,j].sum()),'repeated_value_fraction':float(floor[j]),'published_log2CPM_median':float(median[j]),'median_excluding_numeric_literal_symbols':float(drop_median[j]),'repeated_fraction_excluding_numeric_literal_symbols':float(drop_floor[j]),'PC_scores':pc[j,:3].tolist()} for j,sid in enumerate(p.IDS)],'PCA_sourcewide_gene_standardized_variance_fractions':fraction.tolist(),'descriptive_sourcewide_PC_associations':[{'PC':k+1,'nuisance':label,'Spearman_rho':corr(rankdata(pc[:,k]),rankdata(v))} for k in range(3) for label,v in metrics.items()],'descriptive_selected_gene_export_metric_associations':[{'gene':g,'nuisance':label,'Spearman_rho':corr(rankdata(V[i]),rankdata(v))} for g,i in selected.items() for label,v in metrics.items()],'all12_fixed_internal_partialrank_tests':tests,'no_linked_individual_outcomes':True};dest=OUT/'PeerJ-fullmatrix-export-forensics-actual.json';dest.write_text(json.dumps(sm.clean(result),allow_nan=False,default=str));print('EMC_PEERJ_FULLMATRIX_FORENSICS_BEGIN');print(json.dumps(sm.clean(result),allow_nan=False,default=str));print('EMC_PEERJ_FULLMATRIX_FORENSICS_END')
if __name__=='__main__':main()
