import json,hashlib,datetime,urllib.request,gzip,csv,re,zipfile
from pathlib import Path
from openpyxl import load_workbook
import spatial_marker_followthrough_actual as sm
from pediatric_and_ewing_processed_exports_actual import TARGETS
OUT=sm.OUT;ROOTS=[OUT,Path('restored-artifacts'),Path('restored-artifacts-extra'),Path('restored-artifacts-third')]
def rec(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
 return {'saved':str(p),'bytes':p.stat().st_size,'sha256':h.hexdigest()}
def locate(name):return sorted({p for root in ROOTS for p in root.rglob(name)})
def main():
 OUT.mkdir(parents=True,exist_ok=True);ps=locate('exact-RMS-publisher-and-PeerJ-project-metadata-actual.json');assert ps
 d=json.loads(ps[0].read_text());assert d['schema']=='emc-exact-publisher-project-followthrough/1' and d['RMS_publisher']['exact_primary_PIIs']==['S153458072200243X']
 assets=d['RMS_publisher']['asset_attempts'];assert len(assets)==16;records=[]
 for k in range(2,17):
  old=assets[k-1];name=Path(old['receipt']['saved']).name;paths=locate(name);assert paths,'Restore acquired declaredasset '+name
  p=next((v for v in paths if rec(v)['sha256']==old['receipt']['sha256']),None);assert p
  acquisition={'method':'restored frozen acquired source','earlier_receipt':old['receipt']}
  if k==4:
   p=OUT/name;url=old['receipt']['url'];assert url=='https://ars.els-cdn.com/content/image/1-s2.0-S153458072200243X-mmc4.xlsx'
   h=hashlib.sha256();n=0;cap=512*1024*1024
   with urllib.request.urlopen(urllib.request.Request(url,headers={'User-Agent':'public-sarcoma-reanalysis/1','Accept-Encoding':'identity'}),timeout=90) as r,p.open('wb') as f:
    length=r.headers.get('Content-Length');acquisition.update(method='Complete stream ONLY previously32MiB-truncated mmc4',url=url,final_url=r.geturl(),http_status=r.status,Content_Length=length,cap_bytes=cap)
    while True:
     b=r.read(1024*1024)
     if not b:break
     n+=len(b);assert n<=cap,'Asset>512MiB cap; retainedpartial is not missing source';f.write(b);h.update(b)
   if length is not None:assert n==int(length)
   acquisition.update(complete_bytes=n,complete_sha256=h.hexdigest(),prior_truncation_bytes=old['receipt']['bytes'])
  item={'declared_attachment':old['NIHMS_attachment'],'source':rec(p),'acquisition':acquisition,'sheets':[]};records.append(item)
  with p.open('rb') as f:item['magic_hex']=f.read(16).hex()
  try:
   with zipfile.ZipFile(p) as z:
    assert 'xl/workbook.xml' in z.namelist();item['ZIP_members']=len(z.namelist());item['ZIP_total_declared_uncompressed_bytes']=sum(v.file_size for v in z.infolist())
   wb=load_workbook(p,read_only=True,data_only=True,keep_links=False)
   for si,ws in enumerate(wb.worksheets):
    target=OUT/f'RMS-restored-mmc{k}-sheet{si}-complete.tsv.gz';nr=0;nonempty=0;heads=[];models=[];targets=[];allmodels=set();terms=set();barcodes=0;cols=0
    with gzip.open(target,'wt',newline='') as fh:
     writer=csv.writer(fh,delimiter='\t')
     for row in ws.iter_rows(values_only=True):
      nr+=1;writer.writerow(row);cols=max(cols,len(row));s=[str(v) for v in row if v is not None];nonempty+=bool(s);joined='\t'.join(s)
      if nr<=12:heads.append({'row_1based':nr,'first100_cells':list(row[:100]),'cells':len(row)})
      hits=sorted(set(re.findall(r'SJRHB\d+',joined)));allmodels.update(hits);terms.update(m.group().lower() for m in re.finditer(r'barcode|cell state|myoblast|paraxial|myocyte|malignant|histolog|patient',joined,re.I))
      barcodes+=sum(bool(re.fullmatch(r'[ACGT]{12,24}(?:-\d+)?',v)) for v in s)
      if hits and len(models)<60:models.append({'row_1based':nr,'root_mentions':hits,'first100_cells':list(row[:100])})
      gs=sorted(set(s)&set(TARGETS))
      if gs and len(targets)<200:targets.append({'row_1based':nr,'literal_target_strings':gs,'first100_cells':list(row[:100])})
    item['sheets'].append({'sheet':ws.title,'sheet_state':ws.sheet_state,'rows_streamed':nr,'max_cells_per_row':cols,'nonempty_rows':nonempty,'all_source_values':rec(target),'first12_rows':heads,'all_literal_roots':sorted(allmodels),'first60_root_rows':models,'first200_target_rows':targets,'annotation_terms':sorted(terms),'literal_barcode_like_cells':barcodes})
   wb.close();item['status']='complete_all_sheet_source_reparse'
  except Exception as e:item.update(status='parser_error_not_source_absence',error_type=type(e).__name__,error=str(e))
 pack=OUT/'RMS-complete-publisher-sheets-TSV.zip'
 with zipfile.ZipFile(pack,'w',compression=zipfile.ZIP_STORED) as z:
  for path in sorted(OUT.glob('RMS-restored-mmc*-complete.tsv.gz')):z.write(path,path.name)
 out={'schema':'emc-RMS-restored-publisher-assets/1','executed_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'input_receipt':rec(ps[0]),'all15_declared_XLSX':records,'complete_table_archive':rec(pack),'limits':['Fullsource TSV.gz exported without inferred labels','Only previoustruncatedmmc4 re-fetched; counts/metadatafits not repeated']}
 dest=OUT/'RMS-restored-publisher-assets-actual.json';dest.write_text(json.dumps(sm.clean(out),default=str,allow_nan=False))
 print('EMC_RMS_RESTORED_PUBLISHER_ASSETS_BEGIN');print(json.dumps(sm.clean(out),default=str,allow_nan=False));print('EMC_RMS_RESTORED_PUBLISHER_ASSETS_END')
if __name__=='__main__':main()
