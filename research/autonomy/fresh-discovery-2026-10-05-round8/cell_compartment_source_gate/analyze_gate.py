#!/usr/bin/env python3
"""Source/specimen gate only; reproduces exact S1 and metadata, never reads matrices."""
import argparse,collections,datetime,hashlib,io,json,pathlib,re,zipfile,xml.etree.ElementTree as E
from openpyxl import load_workbook
P=pathlib.Path(__file__).resolve().parent
ap=argparse.ArgumentParser();ap.add_argument('--prior-root',default='/workspace/emc-r6-single_cell/research/autonomy/fresh-discovery-2026-10-04-round6/single_cell');ap.add_argument('--output',default=str(P/'OBSERVATIONS.json'));args=ap.parse_args();Q=pathlib.Path(args.prior_root);C=P/'.cache'
def sha(b):return hashlib.sha256(b).hexdigest()
reuse=json.load(open(P/'REUSED-EVIDENCE.json'))
for x in reuse['records']:assert sha((Q/x['path']).read_bytes())==x['sha256'],x['path']
receipt_checks=[]
for f in sorted(P.glob('*receipts*.json')):
 for x in json.load(open(f)):
  if 'cache_path' in x:
   source=C/pathlib.Path(x['cache_path']).name;assert sha(source.read_bytes())==x['sha256'];receipt_checks.append({'receipt':f.name,'source':source.name,'sha256':x['sha256']})
with zipfile.ZipFile(C/'supplementaryFiles.zip') as z:
 members=[{'path':x.filename,'bytes':x.file_size,'compressed_bytes':x.compress_size} for x in z.infolist()]
 name='ccr-23-2976_supplementary_table_s1_suppts1.xlsx';b=z.read(name);assert b==(C/'Supplementary-Table-S1.xlsx').read_bytes()
 s1=load_workbook(io.BytesIO(b),read_only=True,data_only=True);s=s1['Suppl_Table_1'];rows=list(s.iter_rows(values_only=True));assert len(rows)==11 and len(rows[1])==11
 complete_rows=[dict(zip(rows[1],row)) for row in rows[2:]];ids=[x['Sample'] for x in complete_rows];assert ids==['GS001','S167','S322','S559','S708','S408','S410','S914','S956']
 gs=complete_rows[0];assert gs['cell count post-filter']==2929 and gs['median features post-filter']==648 and gs['Whole Geome Sequencing']=='N' and gs['TCR Sequencing']=='N'
 frozen_sum=sum(x['cell count post-filter'] for x in complete_rows[1:]);assert frozen_sum==75716
 pptname='ccr-23-2976_supplementary_figure_s1_suppfs1.pptx';ppt=z.read(pptname)
 with zipfile.ZipFile(io.BytesIO(ppt)) as zz:
  ppt_text={n:[x.text for x in E.fromstring(zz.read(n)).iter() if x.tag.endswith('}t')] for n in zz.namelist() if n.startswith('ppt/slides/') and n.endswith('.xml')}
 additional_supplement_identity=[]
 for othername in ['ccr-23-2976_supplementary_table_s2_suppts2.xlsx','ccr-23-2976_supplementary_table_s3_suppts3.xlsx']:
  content=z.read(othername);book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
  for sheet in book:
   sheetrows=list(sheet.iter_rows(values_only=True));identityhits=[cell for row in sheetrows for cell in row if isinstance(cell,str) and re.search('GS00[12]|extraskeletal|myxoid chondrosarcoma|NR4A3|diagnosis|histolog',cell,re.I)]
   assert not identityhits
   samples=sorted({row[11] for row in sheetrows[2:] if len(row)>11 and isinstance(row[11],str)})
   additional_supplement_identity.append({'member':othername,'sha256':sha(content),'sheet':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,'header':list(sheetrows[1]),'sample_labels':samples,'fresh_identity_hits':identityhits,'scope':'All string fields screened for fresh case/diagnosis linkage; no pathway/TCR numerical outcome analysis or gene-expression interpretation.'})
primary=E.fromstring((Q/'sources/archival_singlecell_2024_PMC11443197.json').read_bytes());ps=[' '.join(x.itertext()) for x in primary.findall('.//body//p')]
fresh_sentence='we also generated single-cell RNA-seq of another sarcoma specimen directly from fresh tissue (Supplementary Table S1)';assert any(fresh_sentence in x for x in ps)
collection=next(x for x in ps if x.startswith('Fresh and frozen tissue specimens were collected'))
choice=next(x for x in ps if x.startswith('Samples for contrasting scRNA-seq with snRNA-seq'))
code={};selected_code={};tree=json.load(open(C/'github-tree.json'));t={x['path']:x for x in tree['tree']};assert tree['truncated'] is False
for filename,repopath in [('fresh-sarc.ipynb','PreProcessing/01_seurat_analysis-fresh-sarc.ipynb'),('quality-table.ipynb','PreProcessing/1.3_figure1_table1.ipynb'),('github-README.md','README.md')]:
 b=(C/filename).read_bytes();blob=hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest();assert blob==t[repopath]['sha'],repopath
 code[filename]={'repository_path':repopath,'git_blob_sha1':blob,'bytes':len(b),'sha256':sha(b)}
 if filename.endswith('ipynb'):
  n=json.loads(b);hits=[]
  for i,c in enumerate(n['cells']):
   text=''.join(c['source'])
   if re.search('pat_list|sample_type|fresh.merge|idkey|Read10X_h5|readRDS|saveRDS',text):hits.append({'cell':i,'source':text})
  selected_code[filename]=hits
  allsource='\n'.join(''.join(c['source']) for c in n['cells']);assert 'gs001_sarc_scseq' in allsource and 'gs002_sarc_scseq' in allsource
  assert not re.search('extraskeletal|myxoid chondrosarcoma|NR4A3|histology|diagnosis|tumor pathology',allsource,re.I)
# Reuse exact prior library eligibility; no new diagnosis inferred from expression or CNA.
old=json.load(open(Q/'GSM-ELIGIBILITY-OBSERVATIONS.json'));named=old['by_series']['GSE243381'];assert len(named)==21
libraries=[]
for gsm in named:
 x=old['libraries'][gsm];b=(Q/'sources'/f'{gsm}.json').read_bytes();assert sha(b)==x['source_sha256']
 fields={k:x[k] for k in ['accession','source_sha256','title','source_name_ch1','characteristics_ch1','relation','library_strategy'] if k in x};fields['status']='verified evaluation reused';fields['scope']='Published non-EMC identity/condition; no numerical matrix or CNA interpretation.';libraries.append(fields)
series={}
for acc in ['GSE243377','GSE243379','GSE243380','GSE243381']:
 source=Q/'sources/GSE243381.json' if acc=='GSE243381' else C/f'{acc}-self.txt';b=source.read_bytes();lines=b.decode().splitlines();attrs=collections.defaultdict(list)
 for l in lines:
  m=re.match(r'!Series_([^=]+?) = (.*)',l)
  if m:attrs[m[1]].append(m[2])
 series[acc]={'sha256':sha(b),'title':attrs['title'],'sample_id':attrs['sample_id'],'summary':attrs['summary'],'overall_design':attrs['overall_design'],'relation':attrs['relation']}
assert set(series['GSE243380']['sample_id'])==set(named[-8:]);assert len(series['GSE243377']['sample_id'])==8 and len(series['GSE243379']['sample_id'])==5
assert set(series['GSE243377']['sample_id']+series['GSE243379']['sample_id']+series['GSE243380']['sample_id'])==set(named)
assert 'SuperSeries of: GSE243380' in series['GSE243381']['relation'] and 'SubSeries of: GSE243381' in series['GSE243380']['relation']
summ=json.load(open(C/'gds-fresh-label-summary.json'))['result'];unrelated=[{k:summ[u].get(k) for k in ['uid','accession','title','taxon','gse']} for u in summ['uids']];assert len(unrelated)==3 and all(x['taxon']=='Sus scrofa domesticus' for x in unrelated)
project_queries={n:{'count':E.fromstring((C/n).read_bytes()).findtext('Count'),'errors':[x.text for x in E.fromstring((C/n).read_bytes()).findall('.//ErrorList/*')],'limit':'No indexed record returned under these project queries. Not proof of missing controlled-access raw data, not a successful BioSample crosswalk.'} for n in ['biosample-project.xml','sra-project.xml']}
cache_files=[{'path':str(f.relative_to(P)),'bytes':f.stat().st_size,'sha256':sha(f.read_bytes()),'retention':'ignored cache-only original; no matrix staged'} for f in sorted(C.iterdir()) if f.is_file()]
out={'schema':'emc-r8-cell-compartment-source-gate-observations/1','utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'measurement':'Source identity and assay linkage, no CSPG4 expression measured','prospective_plan_sha256':sha((P/'PLAN.json').read_bytes()),'primary_source':{'path':str(Q/'sources/archival_singlecell_2024_PMC11443197.json'),'sha256':sha((Q/'sources/archival_singlecell_2024_PMC11443197.json').read_bytes()),'fresh_source_sentence':fresh_sentence,'collection_quote':collection,'sample_selection_quote':choice},'supplement':{'zip_sha256':sha((C/'supplementaryFiles.zip').read_bytes()),'archive_members':members,'S1_member':name,'S1_sha256':sha((C/'Supplementary-Table-S1.xlsx').read_bytes()),'complete_S1_rows':complete_rows,'printed_fields':99,'frozen_cells_postfilter_sum':frozen_sum,'histology_column_present':False,'S1_figure_member_sha256':sha(ppt),'S1_figure_extracted_slide_text':ppt_text,'additional_supplements_identity_screen':additional_supplement_identity,'figure_limit':'Slide text has assay/comparator labels, no new patient diagnosis; figures and table are published quality comparisons, not CSPG4 values. S2 and S3 all string fields screened for fresh identity/pathology and retain no such labels; numerical outcomes not analyzed. GS001 has no TCR.'},'fresh_conditions':[{'source_label':'GS001','code_alias':'gs001_sarc_scseq','state':'Primary S1 fresh unsorted scRNA-seq','diagnosis':'Unreported in inspected primary/S1/fresh source notebooks','EMC_authentication':'Unresolved, not non-EMC exclusion','GSM':'Not linked in named21/sibling metadata','BioSample':'No source-supported accession crosswalk','patient_condition_overlap':'Not resolved; methods discuss fresh/sequential frozen availability, so do not count as a ninth independent donor.','matrix':'Code refers to local CellBender/Seurat files, no exact public released matrix link for this fresh specimen identified.','status':'pending accessible analysis','barrier':'Identity/pathology plus public assay/cell matrix linkage missing.'},{'source_label':'GS002','code_alias':'gs002_sarc_scseq','state':'Source-code-only additional fresh label and merge directive','diagnosis':'Unreported','EMC_authentication':'Unresolved; source code alone does not prove an actual independent specimen','GSM':'Not linked in named21/sibling metadata','BioSample':'No source-supported accession crosswalk','patient_condition_overlap':'Donor, technical replication, condition and relation to GS001 unresolved; do not count as an authenticated additional donor.','matrix':'Local paths/merge code only; S1 lacks GS002 and displayed quality tables include GS001 only.','status':'pending accessible analysis','barrier':'Actual specimen identity and reported/deposited condition not established.'}],'author_source':{'repository':'https://github.com/IzarLab/sarcoma-sn','tree_sha1':tree['sha'],'full_tree_not_truncated':True,'tree_blob_count':sum(x['type']=='blob' for x in tree['tree']),'files_verified_against_git_blob_sha1':code,'selected_exact_code':selected_code,'data_file_scope':'No per-cell matrix object or GS001/GS002 metadata file is tracked in this inspected repository tree; source code uses local data paths. Other notebooks not reanalyzed; no restricted CNA/structural outcomes inspected.'},'series':series,'series_discrepancy':'All three child metadata summaries say4UPS/4INS; primary and actual prior library source audit say5UPS/3INS. Preserve printed facts; do not rewrite histology to fit counts. GSE243381 is superseries; GSE243380 is snRNA child, not parent.','reused_all21_libraries':libraries,'reused_frozen_units':'Eight specimens from six named donors,5UPS/3INS; X/Y repeated pre/post specimens, not eight independent donors. Unchanged verified prior identity audit reused.','focused_public_label_query':{'GDS_count':3,'actual_returned_identities':unrelated,'decision':'All returned GS001/GS002 alias hits are pig gut GSE277706, not this sarcoma; no false crosswalk. Full underscore labels PhraseNotFound.'},'project_crosswalk_queries':project_queries,'receipt_hash_checks':receipt_checks,'cache_source_inventory':cache_files,'new_retained_raw_bytes':sum(x['bytes'] for x in cache_files),'decision':'SHELVE this source-dependent compartment analysis: actual public source gate remains inconclusive for fresh EMC identity and matrix linkage. No CSPG4 negative or positive finding.','coverage_limit':'Selected source gap substantively advanced with authentic S1 and exact code/source metadata, but not closed as a non-EMC exclusion. All other public EMC cell-resolved/CSPG4 sources and suitable fresh case identity/matrices remain pending before any promotion.'}
pathlib.Path(args.output).write_text(json.dumps(out,indent=2,ensure_ascii=False)+'\n')
print(json.dumps({'S1_rows':len(complete_rows),'fields':99,'named_libraries_reused':len(libraries),'new_raw_bytes':out['new_retained_raw_bytes'],'fresh_identity':'unresolved','expression_analysis':False,'OBSERVATIONS_sha256':sha(pathlib.Path(args.output).read_bytes())},indent=2))
