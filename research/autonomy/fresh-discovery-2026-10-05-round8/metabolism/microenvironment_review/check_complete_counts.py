"""Independent accounting/calibration replay; no target-panel interpretation."""
from collections import defaultdict
import datetime
import hashlib
import json
from pathlib import Path

import numpy as np
import openpyxl
from scipy.stats import spearmanr

OUT = Path(__file__).resolve().parent
WORKER = Path('/workspace/emc-r6-single_cell/research/autonomy/fresh-discovery-2026-10-05-round8/microenvironment')
LANE = WORKER if WORKER.exists() else OUT.parents[1]/'microenvironment'
REPO = OUT.parents[4]


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


receipts_path = LANE/'COMPLETE-TEMPO-SOURCE-RECEIPTS.json'
calibration_path = LANE/'TEMPO-CALIBRATION.json'
assert sha(receipts_path) == '371024c9c9d1e38d638e475ccaf2ec5560dbb2ea7c8212b257b70db3a71c68f9'
assert sha(calibration_path) == 'b99eb34da2715b36426ebc094877b1f6101860904c4a62decdb939b96790470b'
receipts = json.loads(receipts_path.read_text())
calibration = json.loads(calibration_path.read_text())
gate = json.loads((LANE/'TEMPO-SOURCE-GATE.json').read_text())
manifest = LANE/'raw/HWT2.0-manifest.xlsx'
expected_manifest = json.loads((LANE/'MANIFEST-SOURCE-RECEIPT.json').read_text())
assert sha(manifest) == expected_manifest['sha256']
book = openpyxl.load_workbook(manifest, read_only=True, data_only=True)
iterator = book.active.iter_rows(values_only=True)
header = next(iterator)
probes = list(iterator)
book.close()
assert len(probes) == 22537
gene_indices = defaultdict(list)
for j, row in enumerate(probes):
    gene_indices[row[1]].append(j)

assert len(receipts) == len({r['sample'] for r in receipts}) == 12
counts = {}
bindings = []
for r in receipts:
    assert r['complete'] and r['compressed_bytes'] == r['expected_bytes']
    assert r['compressed_md5'] == r['expected_md5']
    assert sum(r['classes'].values()) == r['reads']
    assert r['lengths'] == {'50': r['reads']}
    p = LANE/r['count_output']
    assert sha(p) == r['count_sha256']
    with np.load(p, allow_pickle=False) as z:
        vector = z['counts']
        assert vector.shape == (22537,) and np.issubdtype(vector.dtype, np.integer)
        assert (vector >= 0).all()
        assert int(vector.sum()) == r['classes']['unique_full_probe']
        counts[r['sample']] = vector.copy()
    bindings.append({'path': str(p), 'bytes': p.stat().st_size, 'sha256': sha(p)})
assert sum(r['reads'] for r in receipts) == 81123915
assert sum(r['compressed_bytes'] for r in receipts) == 2704945123

filtered = Path(gate['filtered_source']['path'])
if not filtered.exists():
    filtered = REPO/str(filtered).split('/Rare-cancers/', 1)[1]
assert sha(filtered) == gate['filtered_source']['sha256']
book = openpyxl.load_workbook(filtered, read_only=True, data_only=True)
iterator = book.active.iter_rows(values_only=True)
head = next(iterator)
samples = list(head[1:])
published = {row[0]: np.array(row[1:], dtype=float) for row in iterator}
book.close()
assert len(published) == 9500 and set(samples) == set(counts)
panels = json.loads((LANE/'PLAN.json').read_text())['fixed_panels']
masked = set(sum(panels.values(), []))
genes = sorted((set(published) & set(gene_indices))-masked)
assert len(genes) == calibration['actual_control_genes']
matrix = np.array([[int(counts[s][gene_indices[g]].sum()) for s in samples] for g in genes], dtype=float)
depth = np.array([int(counts[s].sum()) for s in samples], dtype=float)
log_cpm = np.log2(matrix/depth*1e6+.5)
reference = np.array([published[g] for g in genes])
assert np.isfinite(reference).all() and np.isfinite(log_cpm).all()
metrics = []
for j, sample in enumerate(samples):
    rho = float(spearmanr(reference[:, j], log_cpm[:, j]).statistic)
    previous = next(x for x in calibration['per_library'] if x['sample'] == sample)
    assert abs(rho-previous['Spearman_within_library']) < 1e-12
    assert int((matrix[:, j] == 0).sum()) == previous['control_unassigned_zero_gene_counts']
    metrics.append({'sample': sample, 'source_control_Spearman': rho,
                    'meets_frozen_estimator_benchmark_090': rho >= .90})
variance = reference.var(axis=1, ddof=1)
selected = np.flatnonzero(variance >= np.quantile(variance, .75))
cross = [float(spearmanr(reference[j], log_cpm[j]).statistic)
         for j in selected if np.ptp(log_cpm[j]) > 0]
median = float(np.median(cross))
assert len(selected) == calibration['published_top_variance_control_genes']
assert len(cross) == calibration['nonconstant_control_gene_correlations']
assert abs(median-calibration['median_crosssample_Spearman']) < 1e-12
passed = all(x['meets_frozen_estimator_benchmark_090'] for x in metrics) and median >= .80
assert passed == calibration['calibration_passed_frozen_benchmarks'] == False
out = {'utc': datetime.datetime.now(datetime.timezone.utc).isoformat(),
       'question': 'Were all12 source libraries completed and are estimator calibration metrics reproducible?',
       'all12_count_vector_hash_and_class_accounting_pass': True,
       'compressed_bytes_expected_ENA_MD5_all12_match': True,
       'total_reads': 81123915, 'total_compressed_bytes': 2704945123,
       'source_control_genes_after_target_mask': len(genes), 'per_library': metrics,
       'source_control_coverage': {'released_source_rows': 9500,
                                   'present_fixed_target_names_masked': sorted(set(published) & masked),
                                   'unmatched_source_gene_labels_not_imputed': sorted(set(published)-set(gene_indices)),
                                   'interpretation': 'Six numeric source gene labels are unmatched; no alias, date conversion or zero was inferred.'},
       'published_variable_control_genes': len(selected),
       'median_crosssample_control_Spearman': median,
       'frozen_all12_estimator_benchmark_pass': passed,
       'interpretation': 'The raw reads are complete. This exact-probe estimator did not meet its prespecified all-library allocation benchmark. This is not invalid-assay evidence, a missing-gene zero, a negative EMC measurement or exhaustion of valid mapping methods. No fixed target expression or ligand-score interpretation was calculated in this peer replay.',
       'coverage': 'Suitable complete raw data remain pending proper mapping/analysis; this blocks promotion of a ligand claim.',
       'bindings': bindings+[{'path': str(p), 'bytes': p.stat().st_size, 'sha256': sha(p)} for p in
                            [receipts_path, calibration_path, manifest, filtered]]}
(OUT/'COMPLETE-COUNT-CALIBRATION-REVIEW.json').write_text(json.dumps(out, indent=2, allow_nan=False)+'\n')
print(json.dumps({'all12': True, 'reads': out['total_reads'], 'bytes': out['total_compressed_bytes'],
                  'controls': len(genes), 'failed_benchmark_libraries': [r['sample'] for r in metrics if not r['meets_frozen_estimator_benchmark_090']],
                  'median_crosssample_control_Spearman': median}, indent=2))
