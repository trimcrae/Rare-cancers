#!/usr/bin/env python3
"""Read-only independent methylation gate challenge; never reads array values."""
import argparse,collections,datetime,hashlib,json,pathlib,re,subprocess,xml.etree.ElementTree as E
from openpyxl import load_workbook
P=pathlib.Path(__file__).resolve().parent
ap=argparse.ArgumentParser();ap.add_argument('--peer-root',default='/workspace/emc-r6-challenge/research/autonomy/fresh-discovery-2026-10-05-round8/methylation');ap.add_argument('--output',default=str(P/'METHYLATION-INDEPENDENT-CHECK.json'));args=ap.parse_args()
Q=pathlib.Path(args.peer_root)
def sha(f):return hashlib.sha256(f.read_bytes()).hexdigest()
old='research/autonomy/fresh-discovery-2026-10-04-round3/genomics/reassessment/methylation-tables.json'
b=subprocess.check_output(['git','show','80c40ed9fb1af0086bdfc9754c97b176490996e0:'+old],cwd='/workspace/Rare-cancers');a=json.loads(b);table=a['41467_2020_20603_MOESM4_ESM.xlsx']['Tabelle1'];heads=table[0]
refs=[dict(zip(heads,row)) for row in table[1:] if row and str(row[0]).startswith('REFERENCE_SAMPLE')]
soft=(Q/'sources/GSE140686_metadata.txt').read_text();byid={};allgsm={}
for block in soft.split('^SAMPLE = ')[1:]:
 lines=block.splitlines();gsm=lines[0].strip();fields=collections.defaultdict(list)
 for line in lines[1:]:
  m=re.match(r'!Sample_([^=]+?) = (.*)$',line)
  if m:fields[m[1]].append(m[2])
 allgsm[gsm]=dict(fields)
 for desc in fields.get('description',[]):
  if re.match(r'^(REFERENCE_SAMPLE|VALIDATION_SAMPLE) \d+$',desc):
   assert desc not in byid;byid[desc]=(gsm,dict(fields))
controls={'Synovial sarcoma','Myxoid liposarcoma','Low-grade fibromyxoid sarcoma','Sclerosing epithelioid sarcoma'}
emc=[r for r in refs if r['Diagnosis']=='Extraskeletal myxoid chondrosarcoma']
ctl=[r for r in refs if r['Manifestation']=='Primary' and r['DNA']=='FFPE' and r['Diagnosis'] in controls]
assert len(emc)==10 and len(ctl)==62 and len(refs)==1077 and len(allgsm)==1505
actual={};agefields=[]
for r in emc+ctl:
 gsm,meta=byid[r['ID']];plat=meta['platform_id'];assert len(plat)==1
 idat=r['IDAT'];files=meta['supplementary_file'];assert any(idat+'_Grn.idat' in x for x in files);assert any(idat+'_Red.idat' in x for x in files)
 assert meta['description'].count(r['ID'])==1
 ages=[x for k,values in meta.items() if k in ['title','description','characteristics_ch1'] for x in values if re.search(r'\bage\s*[:=]',x,re.I)]
 agefields+=ages;actual[r['ID']]={'GSM':gsm,'platform':plat[0],'slide':idat.split('_')[0],'supplier':r['Supplier'],'batch':r['Batch'],'diagnosis':r['Diagnosis'],'IDAT':idat,'age_annotations':ages,'both_original_IDAT_files_match':True}
checks=[]
for r in emc:
 e=actual[r['ID']];sameplat=[c for c in ctl if actual[c['ID']]['platform']==e['platform']];same_source=[c['ID'] for c in sameplat if c['Supplier']==e['supplier'] and c['Batch']==e['batch']];same_slide=[c['ID'] for c in sameplat if actual[c['ID']]['slide']==e['slide']]
 checks.append({'EMC_ID':r['ID'],'GSM':e['GSM'],'IDAT':e['IDAT'],'platform':e['platform'],'same_platform_controls':len(sameplat),'same_supplier_and_batch_controls':same_source,'same_slide_controls':same_slide,'individual_age_released':bool(e['age_annotations'])})
owner=json.load(open(Q/'EMC-CONTROL-METADATA-AUDIT.json'));assert checks==owner['matching_feasibility'];assert not agefields
assert sum(bool(x['same_slide_controls']) for x in checks)==7
slidepairs={x['EMC_ID']:x['same_slide_controls'] for x in checks};unmatched=[x['EMC_ID'] for x in checks if not x['same_supplier_and_batch_controls'] and not x['same_slide_controls']]
assert len(unmatched)==3
# Independent case21 reading from the original workbook and actual GEO metadata.
w=load_workbook(Q/'sources/myoepithelial2024-S2.xlsx',read_only=True,data_only=True);case21=[];ambiguous=[]
for s in w:
 rows=list(s.iter_rows(values_only=True))
 for row in rows:
  if row and row[0]=='Case 21':case21.append({'sheet':s.title,'cells':list(row)})
  if row and row[0] in ['Case 19','Case 30']:ambiguous.append({'sheet':s.title,'cells':list(row)})
assert len(case21)==1;row=case21[0]['cells'];assert row[1:4]==[35,'M','Anterior abdominal wall'];assert row[7]=='Yes' and row[20]=='TAF15::NR4A3'
metsoft=(Q/'sources/GSE279837_metadata.txt').read_text();samples=metsoft.split('^SAMPLE = ')[1:];assert len(samples)==30
hit=[x for x in samples if '!Sample_title = Case 21\n' in x];assert len(hit)==1
assert hit[0].startswith('GSM8581809\n') and '203135920060_R06C01_Grn.idat' in hit[0] and 'GPL21145' in hit[0]
# Method evidence inspected directly; extract only discriminating passages.
methods={}
for name in ['epitoc2_primary.xml','improved_counter_2024.xml']:
 f=Q/'sources'/name;r=E.fromstring(f.read_bytes());ps=[' '.join(x.itertext()) for x in r.findall('.//body//p')]
 if name.startswith('epitoc2'):
  selectors=['This resulted in 163 CpGs','If the ages of the samples are known','In total, we considered 16 cancer types']
 else:selectors=['We analyzed 32 cancer types','different cell types in a bulk-sample may have different mitotic ages','371 CpGs','the observed strong correlation between stemTOC']
 selected=[]
 for key in selectors:
  hits=[x for x in ps if key in x];assert hits,(name,key);selected.append(hits[0])
 methods[name]={'bytes':f.stat().st_size,'sha256':sha(f),'selected_exact_primary_passages':selected}
assert 'SARC' in methods['improved_counter_2024.xml']['selected_exact_primary_passages'][0]
# Independent omission screen of retained literature records, not a fresh exhaustive search.
lit={}
for f in (Q/'sources').glob('*literature*.json'):
 for x in json.load(open(f)).get('resultList',{}).get('result',[]):lit[(x.get('source'),x['id'])]=x
potential=[{k:x.get(k) for k in ['id','pmcid','doi','title']} for x in lit.values() if re.search('extraskeletal.myxoid|NR4A3|myxoid chondrosarcoma',x.get('title','')+' '+x.get('abstractText',''),re.I) and re.search('methylat|epigenetic|clock',x.get('title','')+' '+x.get('abstractText',''),re.I)]
freeze=json.load(open(Q/'PILOT-FREEZE.json'));bound=[]
for item in freeze['bindings']:
 f=Q/item['path'];bound.append({'path':item['path'],'actual_sha256':sha(f),'matches_original_pilot_freeze':sha(f)==item['sha256']})
sciencefiles=['REPORT.txt','DECISION.json','COVERAGE.json','METHOD-APPLICABILITY.json','EMC-CONTROL-METADATA-AUDIT.json','ADDITIONAL-COLLECTION-OBSERVATIONS.json','VALIDATION.json','MANIFEST.json','PILOT-CONTRACT.json','PILOT-FREEZE.json']
amendment_names=['AMENDMENT-01.json','AMENDMENT-02-STEMTOC-REASSESSMENT.json']
amendments={name:json.load(open(Q/name)) for name in amendment_names}
amendment_hash_checks=[]
for name,record in amendments.items():
 for item in record.get('original_bindings',[])+record.get('bindings',[]):
  assert sha(Q/item['path'])==item['sha256'],(name,item['path'])
  amendment_hash_checks.append({'amendment':name,'path':item['path'],'actual_sha256':sha(Q/item['path']),'matches':True})
assert amendments['AMENDMENT-01.json']['primary_binding']['sha256']==sha(Q/'sources/improved_counter_2024.xml')
assert amendments['AMENDMENT-01.json']['source_quote'] in methods['improved_counter_2024.xml']['selected_exact_primary_passages'][0]
assert 'not mathematically required' in amendments['AMENDMENT-02-STEMTOC-REASSESSMENT.json']['age_distinction']
assert 'No numerical methylation values' in amendments['AMENDMENT-02-STEMTOC-REASSESSMENT.json']['before_outcomes']
sciencefiles+=amendment_names
review={'schema':'emc-r8-independent-methylation-review/1','utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'reviewer':'Codex /root/radiotherapy','reviewed_packet':str(Q),'reviewed_files':{x:{'sha256':sha(Q/x),'bytes':(Q/x).stat().st_size} for x in sciencefiles},'scope':'Independent source/IDAT/case/selection/method/novelty/coverage challenge, no full methylation values or array reanalysis. Read-only peer inputs; own adaptation output only.','source_bindings':{'prior_R3_derived_table':{'path':old,'git_revision':'80c40ed9fb1af0086bdfc9754c97b176490996e0','sha256':hashlib.sha256(b).hexdigest()},'GSE140686_metadata':{'sha256':sha(Q/'sources/GSE140686_metadata.txt')},'myoepithelial_S2':{'sha256':sha(Q/'sources/myoepithelial2024-S2.xlsx')},'GSE279837_metadata':{'sha256':sha(Q/'sources/GSE279837_metadata.txt')},'GSE243075_series':{'sha256':sha(Q/'sources/GSE243075_series.txt')},**methods},'metadata_check':{'reference_rows':len(refs),'all_GSM_records':len(allgsm),'EMC':len(emc),'controls':len(ctl),'control_diagnoses':dict(collections.Counter(r['Diagnosis'] for r in ctl)),'age_metadata_available_in_selected72':0,'primary_FFPE_all72':all(r['Manifestation']=='Primary' and r['DNA']=='FFPE' for r in emc+ctl),'actual_original_IDAT_pairs_checked':len(actual),'EMC_platforms':dict(collections.Counter(actual[r['ID']]['platform'] for r in emc)),'control_platforms':dict(collections.Counter(actual[r['ID']]['platform'] for r in ctl)),'EMC_supplier_counts':dict(collections.Counter(r['Supplier'] for r in emc)),'matching_checks_identical_to_owner':True,'all10_matching_checks':checks,'unmatched_EMC':unmatched,'source_pairing_limit':'Same platform/slide/supplier is a technical pairing, not donor linkage or independent scientific replication. Sclerosing-epithelioid-sarcoma printed diagnosis maps to the reported SEF methylation class; no silent recoding to another histology.'},'case21':{'original_S2':case21,'GEO_case21_count':len(hit),'GSE279837_total_conditions':30,'GSM':'GSM8581809','IDAT':'203135920060_R06C01_Grn.idat','age':35,'sex':'M','mitotic_activity_gt5_per10_HPF':True,'fusion':'TAF15::NR4A3','identity_status':'Prior R2 verified case21 identity/reclassification reused, not new finding. New age metadata checked here.','binding_limit':'Study title/design/case label support GSE279837 association, but manuscript names unrelated GSE243075; source-level association is not a proved exact manuscript/IDAT donor crosswalk.','case19_case30_limit':'Their original EMCS-like clustering/negativeRNA and ambiguous morphology do not authenticate additional EMC.'},'method_value_challenge':{'epiTOC2':'163-CpG cumulative TNSC can be calculated without age; age-adjusted rate R=TNSC/age requires age. Primary validation is not EMC/sarcoma-specific. No categorical claim of sarcoma inapplicability is justified.','stemTOC':'371-CpG counter has normal mesenchymal cell-line evidence; original data methods explicitly include32TCGA types including SARC. The15-type subset is a restricted endpoint, not the entire cancer dataset. This is stronger general plausibility, not authenticated EMC validation.','focused_clarification':'Verified and accepted additive AMENDMENT-01:32types including SARC versus15subset; initial frozen method/report/decision bytes preserved. AMENDMENT-02 separately reassesses the stronger method before outcomes rather than inheriting epiTOC2 rate limits.','why_no_worthwhile_replicated_contrast':'No score or difference was measured. stemTOC cumulative scoring does not require individual age, and favorable mesenchymal/SARC evidence supports plausibility. Nevertheless no independently measured EMC state endpoint or equivalent independent replication is released: one LGFMS control shared across six EMC and seven MLS for another EMC couple histology to source, three EMC unmatched. One aged Case21 does not supply cohort replication. Bulk cell-history weights and unknown probe/normalization transport remain material. A new classifier-separated descriptor alone would not meet the frozen useful state question. Missing age blocks epiTOC2 rate and limits interpretation of cumulative exposure, not computation of stemTOC.','no_matrix_negative':'Not downloading/streaming the5.29GB inputs is a staged decision, not a negative methylation result, an access failure or evidence of low proliferation. Actual accessible values and suitable unexplored identities remain pending and block advancement.'},'independent_omission_check':{'cached_literature_union':len(lit),'gene_disease_epigenetic_metadata_candidates':potential,'meaning':'Independent retained-metadata eligibility challenge only. Broad model/cohort/classifier and numerical arrays remain incomplete; no exhaustive public-EMC coverage certification.'},'original_pilot_hash_checks':bound,'amendment_hash_checks':amendment_hash_checks,'method_specific_value_reassessment':'AMENDMENT-02 accepted after direct source/metadata challenge: the stronger cumulative method is plausibly computable and missing age is not its standalone defeat. Shelving rests on lack of an independently interpretable, replicated useful EMC state contrast, histology/source coupling and unresolved mixtures. All numerical matrices remain pending accessible analysis, so no measured negative, complete coverage or exhaustion is claimed.','prior_review_snapshot':{'path':'METHYLATION-INDEPENDENT-CHECK-INITIAL.json','sha256':sha(P/'METHYLATION-INDEPENDENT-CHECK-INITIAL.json'),'status':'Preserved pre-amendment review; final receipt supersedes requested clarification status.'},'decision':'Shelving supported after favorable method correction and prospective method-specific value reassessment. No new publication-worthy finding or methylation negative demonstrated.','live_owned_processes':0}
out=pathlib.Path(args.output);out.write_text(json.dumps(review,indent=2,ensure_ascii=False)+'\n')
print(json.dumps({'EMC':len(emc),'controls':len(ctl),'metadata_GSM':len(allgsm),'matched_checks':True,'agefields':len(agefields),'actualIDATpairs':len(actual),'case21_age':35,'source_mismatch_count':sum(not x['matches_original_pilot_freeze'] for x in bound),'review_sha256':sha(out)},indent=2))
