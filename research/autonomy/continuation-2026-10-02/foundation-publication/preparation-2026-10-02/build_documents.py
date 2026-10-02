"""Render the frozen Markdown sources to editable Word files; no science execution."""
import re, pathlib, hashlib, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
HERE=pathlib.Path(__file__).resolve().parent

def inline(p,text):
    pattern=r'(\[[^\]]+\]\(https?://[^)]+\)|https?://[^\s]+|\*\*[^*]+\*\*|`[^`]+`)'
    for token in re.split(pattern,text):
        if not token: continue
        m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',token)
        if m or token.startswith('http'):
            label,url=m.groups() if m else (token.rstrip('.,'),token.rstrip('.,'))
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'000000');pr.append(c);run.append(pr)
            t=OxmlElement('w:t');t.text=label;run.append(t);link.append(run);p._p.append(link)
            if not m and token!=label:p.add_run(token[len(label):])
        else:
            r=p.add_run(token.strip('*') if token.startswith('**') else token.strip('`') if token.startswith('`') else token)
            if token.startswith('**'):r.bold=True
            if token.startswith('`'):r.font.name='Courier New';r.font.size=Pt(10)

def build(name):
    md=(HERE/(name+'.md')).read_text(encoding='utf-8')
    if md.startswith('---'):md=md.split('---',2)[2].strip()
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(0.8);sec.left_margin=sec.right_margin=Inches(0.8)
    for style in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        s=doc.styles[style];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
        fonts=s.element.get_or_add_rPr().find(qn('w:rFonts'))
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
            fonts.attrib.pop(qn('w:'+attr),None)
        s.font.size=Pt(11)
        s.paragraph_format.space_after=Pt(6)
        s.paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(17)
    for heading in ['Heading 1','Heading 2','Heading 3']:
        doc.styles[heading].font.size=Pt(12);doc.styles[heading].font.bold=True
        doc.styles[heading].paragraph_format.space_before=Pt(9)
        doc.styles[heading].paragraph_format.keep_with_next=True
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    # The bundled default template can carry a blue title rule: remove inherited borders.
    for border in list(doc.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    lines=md.splitlines();i=0;plain=[]
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
            cells=[[c.strip() for c in x.strip('|').split('|')] for x in rows if not re.fullmatch(r'[|:\-\s]+',x)]
            table=doc.add_table(rows=len(cells),cols=len(cells[0]));table.autofit=False
            widths=[1.0,0.9,2.75,2.25]
            for ri,row in enumerate(cells):
                trPr=table.rows[ri]._tr.get_or_add_trPr()
                no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
                if ri==0:
                    repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
                for ci,value in enumerate(row):
                    cell=table.cell(ri,ci);cell.width=Inches(widths[ci]);inline(cell.paragraphs[0],value)
                    cell.paragraphs[0].paragraph_format.space_after=Pt(5)
                    for run in cell.paragraphs[0].runs:run.font.size=Pt(10);run.bold=(ri==0)
                    tcPr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        node=OxmlElement('w:'+side);node.set(qn('w:w'),'70');node.set(qn('w:type'),'dxa');margins.append(node)
                    tcPr.append(margins)
                    if ri==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');tcPr.append(shade)
                    plain.append(value)
        else:
            if line.startswith('# '):p=doc.add_paragraph(style='Title');text=line[2:]
            elif line.startswith('### '):p=doc.add_paragraph(style='Heading 2');text=line[4:]
            elif line.startswith('## '):p=doc.add_paragraph(style='Heading 1');text=line[3:]
            else:p=doc.add_paragraph();text=line
            inline(p,text)
            if line.startswith('**Table'):p.paragraph_format.keep_with_next=True
            if name=='online-resource-1' and text=='Licences and attribution':p.paragraph_format.page_break_before=True
            plain.append(re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',text).replace('**','').replace('`',''))
    doc.core_properties.author='Tristan D. McRae'
    doc.core_properties.title=plain[0]
    dest=HERE/(name+'.docx');doc.save(dest)
    (HERE/(name+'.visible-text.txt')).write_text('\n'.join(plain)+'\n',encoding='utf-8')
    return {'file':dest.name,'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),'source_sha256':hashlib.sha256((HERE/(name+'.md')).read_bytes()).hexdigest()}
if __name__=='__main__':
    results=[build(x) for x in ['human-genetics-correspondence','online-resource-1']]
    (HERE/'document-build.json').write_text(json.dumps(results,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(results))
