"""Metadata-only audit; no classifier execution or performance estimation.

Usage: python audit_denominator.py [directory containing checkpoint03 sources]
Requires existing openpyxl. Downloads nothing. Source-native fields are copied,
not remapped. Only the three newly identified QC exclusions are looked up in SDRF.
"""
from pathlib import Path
import collections, csv, hashlib, json, sys, zipfile
import xml.etree.ElementTree as ET
import openpyxl

OUT = Path(__file__).resolve().parent
PRIOR = Path(sys.argv[1]) if len(sys.argv)>1 else OUT.parent/'checkpoint03-methylation'
SOURCES = {
    OUT/'CJP2-7-350-s001.xlsx': '6e549cfb75cc37f7cc7c0efc9c715917282b376e7f172155fee01f538c69587b',
    OUT/'CJP2-7-350-s003.docx': 'b658f5dfad81235e7be7a5546ef503c489b605573ec35cd496ddef5c4bcafe26',
    PRIOR/'lyskjaer.xml': '31c0987584ecc056960d5e4e6800efd3de8858598769240bb432f1f8b0077ba5',
    PRIOR/'E-MTAB-9875.sdrf.txt': 'c6ffc854493f95d3b1c70bc18e0862abf269ea889e7ca2db153d435c80adc216',
}
for p,h in SOURCES.items():
    assert hashlib.sha256(p.read_bytes()).hexdigest()==h, p
w = openpyxl.load_workbook(OUT/'CJP2-7-350-s001.xlsx',read_only=True,data_only=True)
s=w['Table S1']
headers=[c.value for c in s[3]]
rows=[(i,r) for i,r in enumerate(s.iter_rows(min_row=4,values_only=True),4) if isinstance(r[0],int)]
assert len(rows)==986 and len({r[0] for _,r in rows})==986
counts=dict(collections.Counter(r[9] for _,r in rows))
assert counts=={'No':163,'Yes':820,'FAILED':3}
failed=[(i,r) for i,r in rows if r[9]=='FAILED']
assert {r[0] for _,r in failed}=={847,848,982}
assert all(r[13]==r[14]=='FAILED' and r[15]=='NA' for _,r in failed)
revised=[(i,r) for i,r in rows if r[12]=='Yes']
assert len(revised)==6

def tsv(name,head,data):
    with (OUT/name).open('w',newline='',encoding='utf8') as f:
        writer=csv.writer(f,delimiter='\t');writer.writerow(head);writer.writerows(data)

# Exclude demographic/site/purity and full prediction-vector data from this metadata export.
cols=[0,1,2,3,5,6,7,8,9,11,12,13,14,15]
tsv('source-native-eligibility.tsv',['source_sheet','source_row']+[headers[c] for c in cols],
    [['Table S1',i]+[r[c] for c in cols] for i,r in rows])
tsv('revised-diagnoses.tsv',['source_sheet','source_row']+[headers[c] for c in cols],
    [['Table S1',i]+[r[c] for c in cols] for i,r in revised])
tsv('field-dictionary.tsv',['column','source_header'],
    [[openpyxl.utils.get_column_letter(c+1),headers[c]] for c in cols])

lookup={r[1]:(i,r) for i,r in failed}
found=[]
with (PRIOR/'E-MTAB-9875.sdrf.txt').open(encoding='utf8') as f:
    for line,d in enumerate(csv.DictReader(f,delimiter='\t'),2):
        name=d['Array Data File']
        array=name.removesuffix('_Grn.idat').removesuffix('_Red.idat')
        if array in lookup:
            i,r=lookup[array]
            assert d['Source Name']==d['Characteristics[individual]']==f'Case_{r[0]}'
            found.append([r[0],array,i,line,d['Source Name'],d['Characteristics[individual]'],name])
assert len(found)==6
assert all(sum(row[1]==array for row in found)==2 for array in lookup)
tsv('qc-exclusions-sdrf-evidence.tsv',
    ['case_number','sentrix_id','Table S1 row','SDRF line','Source Name','Characteristics[individual]','Array Data File'],found)

root=ET.parse(PRIOR/'lyskjaer.xml').getroot()
coordinates=[]
for sec in root.findall('.//body//sec'):
    title=''.join(sec.find('title').itertext()) if sec.find('title') is not None else ''
    for k,p in enumerate(sec.findall('p'),1):
        text=''.join(p.itertext())
        if any(x in text for x in ['failed the quality control','n = 820','four main groups']):
            coordinates.append({'section_id':sec.get('id'),'title':title,'direct_p_ordinal':k,'text':text})
(OUT/'primary-source-coordinates.json').write_text(json.dumps(coordinates,indent=2,ensure_ascii=False),encoding='utf8')
with zipfile.ZipFile(OUT/'CJP2-7-350-s003.docx') as z:
    doc=ET.fromstring(z.read('word/document.xml'))
    ns='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    paras=[''.join(t.text or '' for t in p.iter(ns+'t')) for p in doc.iter(ns+'p')]
    (OUT/'methods-text.txt').write_text('\n'.join(f'p{i}: {t}' for i,t in enumerate(paras,1)),encoding='utf8')

result={
 'metadata_only':True,'case_rows':len(rows),'source_native_core_field_counts':counts,
 'analyzed_according_to_native_QC_field':sum(v for k,v in counts.items() if k!='FAILED'),
 'qc_cases':[{'case_number':r[0],'sentrix_id':r[1],'row':i,'diagnosis':r[5],
              'histological_subtype':r[6],'source_failure_marker':r[9],
              'case_specific_failed_metric':'not provided in Table S1 or supplementary methods'} for i,r in failed],
 'revision_cases':[r[0] for _,r in revised],
 'nonblank_initial_without_revision':[
     {'case_number':r[0],'row':i,'initial_field':r[8],'revision_flag':r[12]}
     for i,r in rows if r[8] is not None and r[12]!='Yes'],
 'source_summary_discrepancy':{'Table S1 J4:J989 Yes_count':counts['Yes'],
    'Table S3 D4':w['Table S3']['D4'].value,'Table S3 E4':w['Table S3']['E4'].value},
 'sources':[{'file':p.name,'sha256':h,'bytes':p.stat().st_size} for p,h in SOURCES.items()],
}
(OUT/'audit-results.json').write_text(json.dumps(result,indent=2,ensure_ascii=False),encoding='utf8')
print(json.dumps(result,indent=2,ensure_ascii=False))
