#!/usr/bin/env python3
"""Build editable journal files from the adjacent canonical Markdown sources.

Requires python-docx. PDF export and page inspection are separate validation steps.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent


def inline(paragraph, text):
    for piece in re.split(r'(`[^`]+`)', text):
        run = paragraph.add_run(piece.strip('`'))
        if piece.startswith('`'):
            run.font.name = 'Courier New'
            run.font.size = Pt(9)


def build(stem):
    text = (ROOT/(stem+'.md')).read_text(encoding='utf-8')
    if text.startswith('---\n'):
        text = text.split('---', 2)[2].strip()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    sec.left_margin = sec.right_margin = Inches(0.9)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in [('Title', 16), ('Heading 1', 13), ('Heading 2', 11), ('Heading 3', 11)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                cells = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[: -]+', x) for x in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=1, cols=len(rows[0]))
            table.style = 'Table Grid'
            table.autofit = False
            widths = ([2.0, 1.15, 1.3, 1.25] if stem == 'manuscript' and 'Training support' in rows[0][0]
                      else [2.4, 1.1, 1.1, 1.1])
            for j, row in enumerate(rows):
                cells = table.rows[0].cells if j == 0 else table.add_row().cells
                for k, value in enumerate(row):
                    cells[k].width = Inches(widths[k])
                    p = cells[k].paragraphs[0]
                    inline(p, value)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.line_spacing = 1
                    if j == 0: p.runs[0].bold = True
                    for run in p.runs: run.font.size = Pt(10)
                pr = table.rows[j]._tr.get_or_add_trPr()
                pr.append(OxmlElement('w:cantSplit'))
                if j == 0: pr.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph()
            continue
        if line.startswith('# '):
            p = doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('### '):
            p = doc.add_paragraph(line[4:], 'Heading 2')
        elif line.startswith('## '):
            p = doc.add_paragraph(line[3:], 'Heading 1')
        else:
            p = doc.add_paragraph()
            inline(p, line)
            if line.startswith('Table '):
                p.paragraph_format.keep_with_next = True
                for run in p.runs: run.font.size = Pt(10)
        i += 1
    doc.core_properties.author = 'Tristan D. McRae'
    doc.core_properties.title = next(x[2:] for x in lines if x.startswith('# '))
    doc.core_properties.subject = 'Methylation empirical methods manuscript preparation'
    doc.save(ROOT/(stem+'.docx'))
    print(stem+'.docx')


if __name__ == '__main__':
    for name in ['manuscript', 'online-resource-1', 'cover-letter']:
        build(name)
