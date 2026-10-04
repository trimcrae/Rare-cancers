"""Explicit deposited specimen/FISH crosswalk; descriptive identity sensitivity."""
import concurrent.futures
import hashlib
import json
from pathlib import Path
import statistics
import sys
import urllib.request
import xml.etree.ElementTree as ET
BASE=Path(__file__).resolve().parent
sys.path.insert(0,str(BASE.parents[2]/'.cache/python-deps'))
import openpyxl

p=BASE/'peerj-source-s009.xlsx'
assert hashlib.sha256(p.read_bytes()).hexdigest()=='20165fd3ff09ec2d5a24b3c20b78515f42a3309119f248ed055c7484deb45e75'
w=openpyxl.load_workbook(p,read_only=True,data_only=True)
s=w['EMC_Gene-expression_Log2CPM']
ids=list(next(s.iter_rows(min_row=1,max_row=1,values_only=True)))[1:]
rows=[r for r in s.iter_rows(values_only=True) if r[0]=='TMEM266']
assert len(rows)==1
values=dict(zip(ids,rows[0][1:]))
assert len(values)==12

def fetch(url):
    with urllib.request.urlopen(url,timeout=30) as r:
        b=r.read(500001)
    assert len(b)<=500000
    return b,{'url':url,'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest()}

link='https://www.ebi.ac.uk/ena/portal/api/filereport?accession=PRJNA1357027&result=read_run&fields=run_accession,sample_accession,sample_alias,sample_title,experiment_accession&format=json'
b,link_receipt=fetch(link)
run_rows=json.loads(b)
assert {x['sample_alias'] for x in run_rows}==set(values)

def get(i):
    accession='SAMN'+str(i)
    url='https://www.ebi.ac.uk/ena/browser/api/xml/'+accession
    b,receipt=fetch(url)
    r=ET.fromstring(b)
    sample=r.find('SAMPLE')
    attrs={x.findtext('TAG'):x.findtext('VALUE') for x in r.findall('.//SAMPLE_ATTRIBUTE')}
    return {'biosample':accession,'id':sample.attrib['alias'],'FISH_1':attrs['FISH_1'],
            'receipt':receipt,'attributes':attrs}

with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
    meta=list(pool.map(get,range(53073757,53073769)))
assert set(x['id'] for x in meta)==set(values)
assert sum(x['FISH_1']=='EWSR1+' for x in meta)==8
assert sum(x['FISH_1']=='EWSR1-' for x in meta)==4
for x in meta:
    x['TMEM266_Log2CPM']=values[x['id']]
summary={}
for status in ['EWSR1+','EWSR1-']:
    group=[x for x in meta if x['FISH_1']==status]
    v=[x['TMEM266_Log2CPM'] for x in group]
    summary[status]={'n':len(v),'ids':[x['id'] for x in group],
        'median':statistics.median(v),'minimum':min(v),'maximum':max(v)}
out={'scope':'Descriptive EWSR1 break-apart sensitivity; no detection threshold, fusion-partner, localization or group-biology inference',
     'summary':summary,'rows':meta,'run_linkage':run_rows,'run_linkage_receipt':link_receipt,
     'provenance_qualification':'Deposited collection dates/institution metadata are not fully reconciled to manuscript cohort dates/origins; do not infer a new cohort origin from these fields.'}
(BASE/'fish-sensitivity-results.json').write_text(json.dumps(out,indent=2)+'\n')
print(json.dumps(summary))
