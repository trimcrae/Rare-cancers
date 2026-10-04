"""Read-only primary-source eligibility audit. No protein ranking or inference."""
from pathlib import Path
import json, hashlib, xml.etree.ElementTree as ET
from collections import Counter
import openpyxl
P=Path(__file__).resolve().parent
out={}
def digest(name): return hashlib.sha256((P/name).read_bytes()).hexdigest()
def text(e): return ' '.join(' '.join(e.itertext()).split())
for name in ['h19-2026.xml','ngospatial2025.xml','burns2023.xml','tang2024.xml','sarquarium2024.xml']:
 r=ET.parse(P/name).getroot()
 out[name]={'sha256':digest(name),'doi':[e.text for e in r.findall('.//article-id') if e.get('pub-id-type')=='doi'],
 'EMC_relevant_paragraphs':[text(e) for e in r.findall('.//p') if any(k in text(e).lower() for k in ['extraskeletal','emcs','myxoid chondrosarcoma'])],
 'supplement_count':len(r.findall('.//supplementary-material'))}
r=ET.parse(P/'h19-2026.xml').getroot()
out['H19']={'status':'unavailable evidence: EMC-specific tissue values and identities; source evaluated',
 'tables':[text(e) for e in r.findall('.//table-wrap')],
 'data_availability':[text(e) for e in r.findall('.//sec') if any('Data Availability' in text(t) for t in e.findall('title'))],
 'tissue_interpretation':'Two described TMAs include EMC, but final n=150 cohort pools rare histologies in Others n=26. No crosswalk, per-EMC H19 score, denominator, outcome, or independence across TMA sets recovered. Do not assume all 26 are EMC.',
 'perturbation':'Only SW872 and SW982 undergo Gapmer experiments; no authentic EMC perturbation.',
 'model_identity':'Article calls MUG-EMCS extraskeletal myxoid chondrosarcoma; creator resource explicitly calls MUG EMCS extraskeletal mesenchymal chondrosarcoma. No NR4A3 authentication recovered. Excluded from authentic EMC biology pending identity resolution.',
 'creator_source':'https://www.medunigraz.at/en/team-beate-rinner','creator_sha256':digest('mug-creator-resource.html'),
 'decision':'Shelve standalone H19 EMC hypothesis. No EMC-specific empirical novelty or functional inference established.'}
assert 'Extraskeletal mesenchymal chondrosarcoma' in (P/'mug-creator-resource.html').read_text(encoding='utf8')
w=openpyxl.load_workbook(P/'sarquarium2024-EV1.xlsx',read_only=True,data_only=True)
out['Sarquarium']={'sha256':digest('sarquarium2024-EV1.xlsx'),'sheets':[(s.title,s.max_row,s.max_column) for s in w]}
s=w['Culture conditions']
rows=[[v for v in r] for r in s.values if r[0] is not None]
out['Sarquarium']['all_model_rows']=rows
assert len(rows)==18
out['Sarquarium']['decision']='All model rows inspected. No authentic EMC model; SW1353 is conventional chondrosarcoma, not EMC. The omitted EMC cannot be inferred from a general sarcoma label.'
w=openpyxl.load_workbook(P/'burns2023-S1.xlsx',read_only=True,data_only=True)
rows=list(w.worksheets[0].values)
out['Burns2023']={'sha256':digest('burns2023-S1.xlsx'),'histology_counts':dict(zip(rows[5][3:],rows[6][3:])),
 'source_total':rows[6][2], 'decision':'All 11 reported histologies sum to 321, with no residual Other category or EMC. Demonstrably unsuitable for EMC biology; no protein values analyzed.'}
assert sum(out['Burns2023']['histology_counts'].values())==321
if (P/'tang2024-S1.xlsx').exists():
 w=openpyxl.load_workbook(P/'tang2024-S1.xlsx',read_only=True,data_only=True)
 out['Tang2024']={'sha256':digest('tang2024-S1.xlsx'),'sheets':[(s.title,s.max_row,s.max_column) for s in w]}
 # Only metadata header/annotation sheet, never proteomic measurement values.
 for s in w:
  if any(x in s.title.lower() for x in ['clinical','patient','1a','1b']):
   rows=list(s.values)
   out['Tang2024'][s.title]=rows
 rows=out['Tang2024']['Clinical_information']
 assert len(rows)==273
 out['Tang2024']['subtype_column_counts']=dict(Counter(r[7] for r in rows[1:]))
 out['Tang2024']['all_otherFS']=[{'case':r[0],'specific_histology':r[6]} for r in rows[1:] if r[7]=='otherFS']
 out['Tang2024']['interpretation']='All272 clinical rows evaluated. No EMC label or NR4A3 authentication. Seven otherFS rows specify adult fibrosarcoma(4), inflammatory myofibroblastic tumor(2), low-grade myofibroblastoma(1). SS-5 is otherFS with specific histology NA and remains unresolved, not proof of EMC absence. SS-81 has a shifted row: his_subtype1 synovial sarcoma, his_subtype2 SS, his_abbr yes; preserve source without silent correction. No protein values analyzed.'
out['ProCan2021']={'source':'procan2021-poster.pdf','sha256':digest('procan2021-poster.pdf'),
 'evaluated':'Poster title 205; methods 203 samples from178 patients, >30 subtypes. Broad chondrosarcoma category and rare types <5 removed from displayed analyses. No complete specimen labels or EMC-specific values recovered.',
 'status':'unavailable evidence / unresolved eligibility, not excluded as EMC-negative',
 'consequence':'Cannot authenticate or analyze possible EMC in rare unlabelled cases; no absence-of-public-proteomics claim.'}
out['PanAtlas2025']={'source':'pride-pan-atlas.json','sha256':digest('pride-pan-atlas.json'),
 'status':'pending accessible analysis of sample-level eligibility',
 'evaluated':'Project PXD054790 metadata confirms 999 primary tumors,22 types,1129 total samples; not an evaluated EMC specimen cohort. Publisher full page403 on this pass; no specimen crosswalk inspected.',
 'consequence':'Do not exclude possible EMC or claim complete pan-cancer coverage. Staged gap after scientific source gate, not biological negative.'}
out['PRIDE_search']={'extraskeletal_projects':[{'accession':x['accession'],'title':x['title']} for x in json.loads((P/'pride-extraskeletal.json').read_text())],
 'NR4A3_projects':json.loads((P/'pride-nr4a3.json').read_text()),
 'interpretation':'Both extraskeletal hits concern calciprotein particles, not EMC. Metadata-search absence is not absence of unlabelled eligible specimens.'}
(P/'protein-spatial-eligibility.json').write_text(json.dumps(out,indent=2,default=str),encoding='utf8')
print('Saved source eligibility:',', '.join(out))
print('Sarquarium models',len(out['Sarquarium']['all_model_rows'])-1)
if 'Tang2024' in out:print('Tang sheets',out['Tang2024']['sheets'])
