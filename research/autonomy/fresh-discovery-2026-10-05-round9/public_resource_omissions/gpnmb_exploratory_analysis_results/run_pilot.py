"""One explicitly authorized GPNMB pilot; no other target or source retrieval."""
from pathlib import Path
import csv,gzip,hashlib,json,datetime,statistics,zipfile,io,math
import numpy as np
import openpyxl
P=Path(__file__).resolve().parent
C=P.parent/'gpnmb_exploratory_analysis_contract'
ROOT=Path('/workspace/Rare-cancers')
SC=Path('/workspace/emc-r6-single_cell/research/autonomy/fresh-discovery-2026-10-05-round9/rna_processing/gpnmb_all_condition_readiness')
EMC='Extraskeletal myxoid chondrosarcoma';LG='Low-grade fibromyxoid sarcoma';ML='Myxoid liposarcoma';SS='Synovial sarcoma'
def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1<<20),b''):h.update(b)
 return h.hexdigest()
def dump(n,d):(P/n).write_text(json.dumps(d,indent=2,allow_nan=False)+'\n')
AUTH=ROOT/'research/autonomy/fresh-discovery-2026-10-05-round9/GPNMB-NUMERICAL-PILOT-AUTHORIZATION.json'
assert sha(AUTH)=='6e6bb8d2a150193048e53b87bf215b447d6994b07ea0c7c4a6807f9bb696d9ac'
auth=json.loads(AUTH.read_text());deadline=datetime.datetime.fromisoformat(auth['scientific_stage_deadline_utc'])
def guard():assert datetime.datetime.now(datetime.timezone.utc)<deadline,'Scientific deadline reached: no outcome computation allowed'
guard()
for n,h in [('CONTRACT-FROZEN.json',auth['contract_sha256']),('AMENDMENT-01-UNCERTAIN-MLPS-CONTROL.json',auth['AM01_sha256']),('AMENDMENT-02-SIGN-AND-RESAMPLING-INTERPRETATION.json',auth['AM02_sha256'])]:assert sha(C/n)==h,n
locks=json.loads((C/'INPUT-LOCK.json').read_text())['inputs']
for r in locks:
 q=Path(r['path']);assert sha(q)==r['sha256'] and q.stat().st_size==r['bytes'],r['role']
units=json.loads((C/'FROZEN-SPECIMEN-UNITS.json').read_text());sets=units['sets'];hvmeta=units['Hofvander'];bysample={r['sample_id']:r for r in hvmeta}
assert [len(sets[k][EMC]) for k in sets]==[13,12,9]
dump('AUTHORIZATION-AND-INPUT-CHECK.json',{'checked_before_cells_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'authorization_path':str(AUTH),'authorization_sha256':sha(AUTH),'deadline':auth['scientific_stage_deadline_utc'],'locked_inputs':locks,'all_hash_size_checks_pass':True,'contract_sha256':auth['contract_sha256'],'AM01_sha256':auth['AM01_sha256'],'AM02_sha256':auth['AM02_sha256'],'RNG':'numpy.default_rng(20261005) reset per contrast; x choice(10000,nE), then y choice(10000,nC); linear quantiles','no_network':True})
records=[];missing=[]
def number(raw):
 if raw is None or str(raw).strip() in ['', 'NA','NaN','nan','null','None']:return None,'source_missing'
 try:v=float(raw)
 except (ValueError,TypeError):return None,'source_not_numeric'
 return (v,None) if math.isfinite(v) else (None,'source_nonfinite')
def add(ds,feature,sid,raw,unit,hist,ref='',role=''):
 v,reason=number(raw);r={'dataset':ds,'gene':'GPNMB','feature':str(feature),'condition_id':sid,'diagnosis':hist,'unit':unit,'reference':ref,'role':role,'value':v,'status':'evaluated' if reason is None else 'pending_accessible_measurement','missing_reason':reason}
 records.append(r)
 if reason:missing.append({k:r[k] for k in ['dataset','feature','condition_id','missing_reason']})
 return v
def A(x,y):return float(np.mean(x[:,None]<y[None,:])+0.5*np.mean(x[:,None]==y[None,:]))
def distribution(vv):
 good=[x for x in vv if x is not None]
 return {'source_n':len(vv),'finite_n':len(good),'missing_n':len(vv)-len(good),'min':min(good) if good else None,'median':statistics.median(good) if good else None,'max':max(good) if good else None}
def contrast(vals,ids,controls,unit,logRNA=False):
 guard();mx=[s for s in ids if vals.get(s) is None];my=[s for s in controls if vals.get(s) is None]
 if mx or my:return {'status':'fail_closed_incomplete_cells','EMC_ids':ids,'comparator_ids':controls,'missing_EMC':mx,'missing_comparator':my,'unit':unit}
 x=np.array([vals[s] for s in ids],float);y=np.array([vals[s] for s in controls],float);assert len(x) and len(y)
 rng=np.random.default_rng(20261005);xb=rng.choice(x,size=(10000,len(x)),replace=True);yb=rng.choice(y,size=(10000,len(y)),replace=True)
 ab=np.mean(xb[:,:,None]<yb[:,None,:],axis=(1,2))+0.5*np.mean(xb[:,:,None]==yb[:,None,:],axis=(1,2));db=np.median(xb,axis=1)-np.median(yb,axis=1)
 result={'status':'evaluated','EMC_ids':ids,'comparator_ids':controls,'n_EMC':len(x),'n_comparator':len(y),'unit':unit,'EMC':distribution(x.tolist()),'comparator':distribution(y.tolist()),'A_lower':A(x,y),'EMC_minus_comparator_median_shift':float(np.median(x)-np.median(y)),'A_lower_conditional_empirical_resampling_95_range':np.quantile(ab,[.025,.975],method='linear').tolist(),'median_shift_conditional_empirical_resampling_95_range':np.quantile(db,[.025,.975],method='linear').tolist(),'complete_observed_rank_separation':bool(max(x)<min(y) or min(x)>max(y)),'resampling_range_is_not_population_or_selection_adjusted_CI':True,'leave_one_EMC_specimen_out':[{'removed_id':sid,'A_lower':A(np.delete(x,i),y),'median_shift':float(np.median(np.delete(x,i))-np.median(y))} for i,sid in enumerate(ids)]}
 if logRNA:
  xx=np.log2(1+x);yy=np.log2(1+y);delta=np.median(np.log2(1+xb),axis=1)-np.median(np.log2(1+yb),axis=1)
  result['EMC_minus_comparator_median_log2_1plusTPM_shift']=float(np.median(xx)-np.median(yy));result['log2_1plusTPM_shift_conditional_empirical_resampling_95_range']=np.quantile(delta,[.025,.975],method='linear').tolist()
 return result
# Unique symbol, only fixed 59 histotype conditions converted to numbers.
guard();RNA=next(Path(r['path']) for r in locks if r['role']=='Hofvander_RNA');hv={};found=0
with gzip.open(RNA,'rt') as f:
 rd=csv.reader(f,delimiter='\t');head=next(rd);assert len(head[1:])==704 and len(set(head[1:]))==704
 indices={s:head.index(s) for s in bysample}
 for r in rd:
  if r[0]!='GPNMB':continue
  found+=1;assert len(r)==len(head)
  for s,i in indices.items():
   v=add('Hofvander','GPNMB',s,r[i],'TPM',bysample[s]['diagnosis']);assert v is None or v>=0;hv[s]=v
assert found==1 and len(hv)==59
hresults={}
for scope,groups in sets.items():
 hresults[scope]={}
 for hist in [LG,ML,SS]:
  hresults[scope][hist]=contrast(hv,groups[EMC],groups[hist],'TPM',True)
  if hist==ML:hresults[scope][hist+'_uncertain2492omitted']=contrast(hv,groups[EMC],[s for s in groups[hist] if s!='2492-91'],'TPM',True)
 for hist in [LG,ML]:
  variants=[(hist,groups[hist])]+([(hist+'_uncertain2492omitted',[s for s in groups[hist] if s!='2492-91'])] if hist==ML else [])
  for key,controls in variants:
   strata=[];weighted=0.;pairs=0
   for yr in ['2019','2021']:
    es=[s for s in groups[EMC] if bysample[s]['sequencing_year']==yr];cs=[s for s in controls if bysample[s]['sequencing_year']==yr];pp=len(es)*len(cs);aa=None
    if pp and all(hv[s] is not None for s in es+cs):aa=A(np.array([hv[s] for s in es]),np.array([hv[s] for s in cs]));weighted+=pp*aa;pairs+=pp
    strata.append({'year':yr,'EMC_ids':es,'comparator_ids':cs,'pairs':pp,'A_lower':aa})
   hresults[scope][key]['year_point_estimates_only']={'strata':strata,'A_lower_pair_weighted':weighted/pairs if pairs else None,'finite_pairs':pairs,'no_year_interval':True,'not_all_specimen_batch_adjustment':True}
dump('HOFVANDER-RESULTS.json',hresults)
print('PRIMARY_HOFVANDER',json.dumps({k:{'A':v[LG]['A_lower'],'median_shift':v[LG]['EMC_minus_comparator_median_shift'],'range':v[LG]['A_lower_conditional_empirical_resampling_95_range'],'year_A':v[LG]['year_point_estimates_only']['A_lower_pair_weighted']} for k,v in hresults.items()}),flush=True)
# All 42 RMA source observations, one fixed gene-unique cluster.
guard();SOFT=next(Path(r['path']) for r in locks if r['role']=='GSE24369_SOFT');arows=units['ARRAY_SOURCE_RECORDS'];a42=[r for r in arows if r['gse']=='GSE24369'];o16=[r for r in arows if r['gse']=='GSE4303'];assert len(a42)==42 and len(o16)==16
lookup={r['gsm']:r for r in a42};av={};sid=None;table=False
with gzip.open(SOFT,'rt') as f:
 for line in f:
  line=line.rstrip('\r\n')
  if line.startswith('^SAMPLE = '):sid=line.split(' = ',1)[1]
  elif line=='!sample_table_begin':table=True;columns=None
  elif line=='!sample_table_end':table=False
  elif table:
   row=line.split('\t')
   if columns is None:columns=row;continue
   if row[columns.index('ID_REF')]=='8131844':
    assert sid in lookup and sid not in av;av[sid]=add('GSE24369','8131844',sid,row[columns.index('VALUE')],'RMA log2',lookup[sid]['title'])
for s in lookup:
 if s not in av:av[s]=add('GSE24369','8131844',s,None,'RMA log2',lookup[s]['title'])
ae=[r['gsm'] for r in a42 if 'extraskeletal' in r['title'].lower()];ac=[r['gsm'] for r in a42 if 'low-grade' in r['title'].lower()];assert len(ae)==6 and len(ac)==17
array_result={'primary_LGFMS':contrast(av,ae,ac,'RMA log2'),'all42_context_values':av,'context_titles':{r['gsm']:r['title'] for r in a42},'fixed_feature':'8131844','donor_overlap_unresolved':True}
dump('GSE24369-RESULTS.json',array_result);print('PRIMARY_ARRAY',json.dumps(array_result['primary_LGFMS']),flush=True)
# Three gene-unique historical probes, each source reference kept distinct.
guard();ZIP=next(Path(r['path']) for r in locks if r['role']=='original_array_zip');ov={p:{} for p in ['5535','10100','19562']};oldlookup={r['gsm']:r for r in o16}
with zipfile.ZipFile(ZIP) as z:
 with z.open('GSE4303-GPL3290-source-matrix.gz') as raw:
  with gzip.open(raw,'rt') as f:
   table=False;head=None
   for line in f:
    line=line.rstrip('\r\n')
    if line=='!series_matrix_table_begin':table=True;continue
    if line=='!series_matrix_table_end':break
    if table:
     rr=next(csv.reader([line],delimiter='\t'))
     if head is None:head=rr[1:];assert set(head)==set(oldlookup);continue
     if rr[0] in ov:
      assert not ov[rr[0]] and len(rr[1:])==len(head)
      for sid,v in zip(head,rr[1:]):ov[rr[0]][sid]=add('GSE4303',rr[0],sid,v,'log2 tumor/reference',oldlookup[sid]['title'],';'.join(oldlookup[sid]['source_ch1']))
oldcomposite={}
for r in o16:
 s=r['gsm'];vv=[ov[p].get(s) for p in ov];value=statistics.median(vv) if all(v is not None for v in vv) else None
 oldcomposite[s]=add('GSE4303','median_fixed_three',s,value,'log2 tumor/reference',r['title'],';'.join(r['source_ch1']))
old={'per_probe':ov,'all_three_finite_median':oldcomposite,'reference_groups':{g:[r['gsm'] for r in o16 if ';'.join(r['source_ch1'])==g] for g in ['CRH-mRNA','CRH','UHR']},'no_primary_or_secondary_replication':'No LGFMS/MLPS. Mixed CRH/UHR contrasts not computed.','all_EMC_conditions':10}
dump('GSE4303-RESULTS.json',old)
# Both fixed coding peaks; no source DE annotation field is read.
guard();S3=next(Path(r['path']) for r in locks if r['role']=='3SEQ_source' and r['path'].endswith('.gz'));s3={p:{} for p in ['10146','10147']};source_anno=json.loads((SC/'GSE28866-ASSAY-AND-PEAK-ANNOTATION.json').read_text());emc3=source_anno['EMC_header_columns']
with gzip.open(S3,'rt') as f:
 reader=csv.DictReader(f,delimiter='\t');fields=reader.fieldnames;allcols=fields[7:];assert len(allcols)==93
 mls=[c for c in allcols if c.startswith('MLPS_')];sss=[c for c in allcols if c.startswith('SS_')];norm=[c for c in allcols if '_normal_' in c];assert [len(emc3),len(mls),len(sss),len(norm)]==[4,6,3,27];selected=emc3+mls+sss+norm
 for r in reader:
  pk=r['peak']
  if pk not in s3:continue
  assert not s3[pk] and r['gene_symbol']=='GPNMB' and r['peak_exon_gene_symbol']=='GPNMB'
  for s in selected:s3[pk][s]=add('GSE28866',pk,s,r[s],'normalized square-root peak density','EMC' if s in emc3 else 'MLPS' if s in mls else 'SS' if s in sss else 'normal source context')
assert all(len(x)==40 for x in s3.values())
composite={}
for s in selected:
 vv=[s3[p][s] for p in s3];composite[s]=add('GSE28866','median_fixed_two',s,statistics.median(vv) if all(v is not None for v in vv) else None,'normalized square-root peak density','EMC' if s in emc3 else 'MLPS' if s in mls else 'SS' if s in sss else 'normal source context')
three={'per_peak':s3,'median_both_finite':composite,'comparisons':{p:{'secondary_MLPS':contrast(v,emc3,mls,'normalized square-root peak density'),'context_SS':contrast(v,emc3,sss,'normalized square-root peak density')} for p,v in {**s3,'median_fixed_two':composite}.items()},'normal_context_values':{s:composite[s] for s in norm},'no_LGFMS':'No same-assay primary comparator','not_linear_gene_counts':True}
dump('GSE28866-RESULTS.json',three)
# Read only known GPNMB row in the published processed sheet.
guard();T=next(Path(r['path']) for r in locks if r['role']=='published_TempO_Log2CPM');wb=openpyxl.load_workbook(T,read_only=True,data_only=True);ws=wb['EMC_Gene-expression_Log2CPM'];header=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))];expected=['Si01','Si02','Si05','Si09','Si10','Si14','Si15','Si16','Si17','Si19','Si20','Si22'];assert header==['symbol']+expected
row=next(ws.iter_rows(min_row=7134,max_row=7134,values_only=True));assert row[0]=='GPNMB' and len(row)==13
tempo={s:add('TempO_published','GPNMB_Log2CPM',s,v,'published Log2CPM','author-labelled EMC') for s,v in zip(expected,row[1:])};wb.close();dump('TEMPO-PROCESSED-CONTEXT.json',{'values':tempo,'distribution':distribution(list(tempo.values())),'conditions':12,'no_LGFMS_MLPS':'Conditional native context only, case/condition/donor linkage and cross-cohort overlap pending','raw_estimator_failure_not_applied':True})
# Compact every eligible condition/feature, including zeros, highs and missing cells.
guard()
with (P/'MEASUREMENTS.tsv').open('w',newline='') as f:
 wr=csv.DictWriter(f,fieldnames=list(records[0]),delimiter='\t',lineterminator='\n');wr.writeheader();wr.writerows(records)
dump('MISSING-AND-COMPLETENESS.json',{'missing':missing,'measurement_feature_rows':len(records),'counts_by_dataset':{ds:len({r['condition_id'] for r in records if r['dataset']==ds}) for ds in sorted({r['dataset'] for r in records})},'all_core_conditions_accounted':True,'finite_completeness_not_inferred_from_header':True})
dump('SCIENTIFIC-STAGE-CLOSED.json',{'closed_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'deadline':auth['scientific_stage_deadline_utc'],'before_deadline':True,'computed_target':'GPNMB only','new_network_or_originals':0,'new_p_q_outputs':0,'no_culture_values':True,'authorization_sha256':sha(AUTH),'gene_outcomes_read_only_after_authorization':True,'after_closure':'Freeze/archive/validation/commit only; no new outcome computation without a distinct integrating-lead authorization.'})
print('FINISHED',datetime.datetime.now(datetime.timezone.utc).isoformat(), 'measurement_rows',len(records),'missing',len(missing),flush=True)
