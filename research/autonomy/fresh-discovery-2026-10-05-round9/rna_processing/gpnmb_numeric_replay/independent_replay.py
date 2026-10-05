#!/usr/bin/env python3
"""Independent fixed-GPNMB replay. No source outcomes are read without --frozen-owner.
Only the exact root-authorized pilot is replayed; no network or new target selection.
"""
import argparse, csv, gzip, hashlib, io, json, math, pathlib, datetime, zipfile
import numpy as np
ROOT=pathlib.Path(__file__).resolve().parent
CONTRACT=pathlib.Path('/workspace/emc-r6-diagnostic/research/autonomy/fresh-discovery-2026-10-05-round9/public_resource_omissions/gpnmb_exploratory_analysis_contract')
READY=ROOT.parent/'gpnmb_all_condition_readiness'

def sha(p): return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()
def load(p): return json.loads(pathlib.Path(p).read_text())
def write(p,d): pathlib.Path(p).write_text(json.dumps(d,indent=2,allow_nan=False)+'\n')
def num(x):
 try:
  v=float(x); return v if math.isfinite(v) else None
 except (ValueError,TypeError): return None

def pair_effect(x,y):
 # Literal independently implemented empirical pair counting, half ties.
 return float(sum((a<b)+.5*(a==b) for a in x for b in y)/(len(x)*len(y)))
def estimate(x,y):
 if not len(x) or not len(y) or any(v is None for v in list(x)+list(y)):
  return {'status':'missing/nonfinite fixed input; comparison unavailable'}
 x=np.asarray(x,dtype=float); y=np.asarray(y,dtype=float)
 rng=np.random.default_rng(20261005)
 bx=rng.choice(x,size=(10000,len(x)),replace=True)
 by=rng.choice(y,size=(10000,len(y)),replace=True)
 # Per-bootstrap counting accumulated without owner code/vectorized 3D expression.
 below=np.zeros(10000); tied=np.zeros(10000)
 for a in range(len(x)):
  for b in range(len(y)):
   below+=bx[:,a]<by[:,b]; tied+=bx[:,a]==by[:,b]
 boot_a=(below+.5*tied)/(len(x)*len(y))
 boot_shift=np.median(bx,axis=1)-np.median(by,axis=1)
 q=lambda z:np.quantile(z,[.025,.975],method='linear').tolist()
 return {'status':'finite','n_EMC':len(x),'n_control':len(y),'pairs':len(x)*len(y),
 'A_lower':pair_effect(x,y),'median_EMC':float(np.median(x)),
 'median_control':float(np.median(y)), 'median_shift':float(np.median(x)-np.median(y)),
 'A_lower_conditional_range':q(boot_a),'median_shift_conditional_range':q(boot_shift)}

def verify_locked_inputs():
 out=[]
 for r in load(CONTRACT/'INPUT-LOCK.json')['inputs']:
  actual=sha(r['path']); assert actual==r['sha256'],(r['path'],actual,r['sha256'])
  out.append({'path':r['path'],'bytes':pathlib.Path(r['path']).stat().st_size,'sha256':actual,'status':'exact'})
 return out

def extract_fixed():
 locked=load(CONTRACT/'INPUT-LOCK.json')['inputs']; p={r['role']:r['path'] for r in locked if not r['role'].startswith('readiness:')}
 units=load(CONTRACT/'FROZEN-SPECIMEN-UNITS.json')
 cells={}; annotation=[]
 # Only the fixed GPNMB row is parsed numerically. Non-target lines discarded.
 with gzip.open(p['Hofvander_RNA'],'rt') as f:
  header=f.readline().rstrip('\r\n').split('\t'); target=None; matches=0
  for line in f:
   sym=line.partition('\t')[0]
   if sym=='GPNMB': target=line.rstrip('\r\n').split('\t'); matches+=1
  assert matches==1
  bylabel=dict(zip(header[1:],target[1:])); cells['Hofvander']=[{**r,'value':num(bylabel[r['sample_id']])} for r in units['Hofvander']]
 # SOFT sample tables: only the fixed probe's VALUE is evaluated.
 array_roster=load(READY/'ARRAY-CONDITION-ROSTERS.json')['all58_source_records']
 wanted={r['gsm']:r for r in array_roster if r['gse']=='GSE24369'}
 values={}; sample=None; table=False; valindex=None
 with gzip.open(p['GSE24369_SOFT'],'rt',errors='strict') as f:
  for line in f:
   if line.startswith('^SAMPLE = '): sample=line.strip().split(' = ',1)[1]; table=False
   elif line.startswith('!sample_table_begin'): table=True; valindex=None
   elif line.startswith('!sample_table_end'): table=False
   elif table and sample in wanted:
    if valindex is None:
     head=line.rstrip('\r\n').split('\t'); valindex=head.index('VALUE')
    elif line.partition('\t')[0]=='8131844': values[sample]=num(line.rstrip('\r\n').split('\t')[valindex])
 cells['GSE24369']=[{**r,'features':{'8131844':values.get(r['gsm'])},'value':values.get(r['gsm'])} for r in array_roster if r['gse']=='GSE24369']
 # Deposited source matrix in immutable ZIP; only exact 3 probe rows parsed.
 z=zipfile.ZipFile(p['original_array_zip']); member='GSE4303-GPL3290-source-matrix.gz'
 with gzip.open(io.BytesIO(z.read(member)),'rt') as f:
  features={}; head=None
  for line in f:
   if line.startswith('"ID_REF"') or line.startswith('ID_REF'):
    head=next(csv.reader([line],delimiter='\t')); continue
   probe=line.partition('\t')[0].strip('"')
   if probe in ('5535','10100','19562'):
    row=next(csv.reader([line],delimiter='\t')); features[probe]=dict(zip(head[1:],row[1:]))
  assert set(features)=={'5535','10100','19562'}
 out=[]
 for r in array_roster:
  if r['gse']!='GSE4303':continue
  vals={pr:num(features[pr].get(r['gsm'])) for pr in ('5535','10100','19562')}
  out.append({**r,'features':vals,'value':float(np.median(list(vals.values()))) if all(v is not None for v in vals.values()) else None})
 cells['GSE4303']=out
 # 3SEQ field7 is explicitly not read/projected. Fixed rows, numeric fields8 onward.
 seqpath=next(r['path'] for r in locked if r['path'].endswith('GSE28866_normalized.txt.gz'))
 with gzip.open(seqpath,'rt') as f:
  head=f.readline().rstrip('\r\n').split('\t'); fixed={}
  for line in f:
   if line.partition('\t')[0] not in ('10146','10147'):continue
   fields=line.rstrip('\r\n').split('\t'); ident={head[i]:fields[i] for i in range(6)}
   assert ident['gene_symbol']=='GPNMB' and ident['peak_exon_gene_symbol']=='GPNMB'
   annotation.append(ident); fixed[fields[0]]=dict(zip(head[7:],fields[7:]))
  assert set(fixed)=={'10146','10147'}
  normals=[h for h in head[7:] if '_normal_' in h.lower()]; assert len(normals)==27
  eligible=units['3SEQ_source_conditions']+units['3SEQ_comparator_conditions']+normals
  assert len(eligible)==40 and len(set(eligible))==40
  cells['GSE28866']=[{'condition':h,'features':{pr:num(fixed[pr][h]) for pr in ('10146','10147')},'value':float(np.median([num(fixed[pr][h]) for pr in ('10146','10147')])) if all(num(fixed[pr][h]) is not None for pr in ('10146','10147')) else None} for h in eligible]
 # Published processed gene row, no rederived counts.
 import openpyxl
 wb=openpyxl.load_workbook(p['published_TempO_Log2CPM'],read_only=True,data_only=True)
 ws=wb['EMC_Gene-expression_Log2CPM']; head=next(ws.iter_rows(min_row=1,max_row=1,values_only=True)); row=next(ws.iter_rows(min_row=7134,max_row=7134,values_only=True)); assert row[0]=='GPNMB'
 cells['published_TempO']=[{'condition':head[i],'value':num(row[i])} for i in range(1,13)];wb.close()
 return cells,annotation

def point(x,y):
 if not x or not y or any(v is None for v in x+y):return {'status':'missing'}
 return {'n_EMC':len(x),'n_control':len(y),'pairs':len(x)*len(y),'A_lower':pair_effect(x,y),'median_shift':float(np.median(x)-np.median(y))}

def contrast(e,c,rna=False):
 x=[r['value'] for r in e]; y=[r['value'] for r in c]
 out={'EMC_ids':[r.get('sample_id',r.get('gsm',r.get('condition'))) for r in e], 'control_ids':[r.get('sample_id',r.get('gsm',r.get('condition'))) for r in c], 'native_scale':estimate(x,y)}
 if rna:
  lx=[float(np.log2(1+v)) if v is not None else None for v in x];ly=[float(np.log2(1+v)) if v is not None else None for v in y]
  out['log2_1_plus_TPM']=estimate(lx,ly)
 out['leave_one_EMC_out']=[{'removed_id':out['EMC_ids'][i],**point(x[:i]+x[i+1:],y)} for i in range(len(x))]
 if rna:
  out['leave_one_EMC_out_log2_1_plus_TPM']=[{'removed_id':out['EMC_ids'][i],**point(lx[:i]+lx[i+1:],ly)} for i in range(len(lx))]
 return out

def calculate(cells):
 out={'Hofvander':{},'GSE24369':{},'GSE28866':{},'contexts':{}}
 h=cells['Hofvander'];e=[r for r in h if r['diagnosis']=='Extraskeletal myxoid chondrosarcoma']
 assert len(e)==13
 e_sets={'all13':e,'primary12':[r for r in e if r['primary_lesion']], 'known_overlap9':[r for r in e if r['primary_lesion'] and not r['known_overlap']]}
 assert [len(v) for v in e_sets.values()]==[13,12,9]
 for name,es in e_sets.items():
  ctrls={'LGFMS':[r for r in h if r['diagnosis']=='Low-grade fibromyxoid sarcoma'], 'MLPS_all14':[r for r in h if r['diagnosis']=='Myxoid liposarcoma'], 'MLPS_uncertain_omitted13':[r for r in h if r['diagnosis']=='Myxoid liposarcoma' and r['sample_id']!='2492-91'], 'SS_context':[r for r in h if r['diagnosis']=='Synovial sarcoma' and (name=='all13' or r['primary_lesion'])]}
  for cname,cs in ctrls.items():
   key=name+'__'+cname; out['Hofvander'][key]=contrast(es,cs,True)
   if cname=='SS_context':continue
   years=[]
   for year in ('2019','2021'):
    ey=[r for r in es if r['sequencing_year']==year];cy=[r for r in cs if r['sequencing_year']==year]
    years.append({'year':year,'EMC_ids':[r['sample_id'] for r in ey],'control_ids':[r['sample_id'] for r in cy],**point([r['value'] for r in ey],[r['value'] for r in cy])})
   valid=[r for r in years if r.get('pairs',0)>0]
   out['Hofvander'][key]['year_point_only']={'strata':years,'pair_weighted_A_lower':sum(r['pairs']*r['A_lower'] for r in valid)/sum(r['pairs'] for r in valid), 'pairs':sum(r['pairs'] for r in valid),'unmatched_EMC_ids':[r['sample_id'] for r in es if r['sequencing_year'] not in ('2019','2021')],'unmatched_control_ids':[r['sample_id'] for r in cs if r['sequencing_year'] not in ('2019','2021')]}
 a=cells['GSE24369']; e=[r for r in a if r['title'].startswith('Extraskeletal myxoid chondrosarcoma')]
 c=[r for r in a if r['title'].startswith('Low grade fibromyxoid sarcoma') or r['title'].startswith('Low-grade fibromyxoid sarcoma')]
 # Source titles checked exactly by caller if different spelling; no generic identity inference.
 assert len(e)==6 and len(c)==17,(len(e),len(c))
 out['GSE24369']['8131844__EMC_vs_LGFMS']=contrast(e,c)
 seq=cells['GSE28866']
 for feature in ('10146','10147','both_peak_median'):
  rec=[{**r,'value':r['value'] if feature=='both_peak_median' else r['features'][feature]} for r in seq]
  e=[r for r in rec if r['condition'].startswith('EMC_')]
  for cp in ('MLPS_','SS_'):
   c=[r for r in rec if r['condition'].startswith(cp)];out['GSE28866'][feature+'__'+cp.rstrip('_')]=contrast(e,c)
 # No control or reference pooling. All context numeric cells remain separately retained.
 for source in ('GSE4303','published_TempO'):
  out['contexts'][source]={'n_conditions':len(cells[source]),'finite_composites':sum(r['value'] is not None for r in cells[source]),'individual_values_retained':'INDEPENDENT-FIXED-CELLS.json; separate native/deposited units; no incompatible control comparison'}
 return out

def main():
 ap=argparse.ArgumentParser(); ap.add_argument('--frozen-owner',required=True);ap.add_argument('--owner-freeze-sha',required=True); a=ap.parse_args()
 assert sha(a.frozen_owner)==a.owner_freeze_sha,'Owner result freeze mismatch; no outcome access'
 started=datetime.datetime.now(datetime.timezone.utc); write(ROOT/'REPLAY-START.json',{'first_frozen_result_access_at':started.isoformat(),'scientific_deadline':(started+datetime.timedelta(minutes=15)).isoformat(),'owner_freeze':a.frozen_owner,'sha256':a.owner_freeze_sha,'network':0,'new_originals':0})
 write(ROOT/'VERIFIED-SOURCE-HASHES.json',verify_locked_inputs())
 cells,annotation=extract_fixed();write(ROOT/'INDEPENDENT-FIXED-CELLS.json',cells);write(ROOT/'INDEPENDENT-3SEQ-ANNOTATION.json',annotation)
 write(ROOT/'INDEPENDENT-ESTIMANDS.json',calculate(cells))
 print('Fixed source cells and prespecified estimands independently replayed; no other target or source DE field read.')

if __name__=='__main__':main()
