"""Descriptive extraction only; all source units/replicate uncertainty retained."""
import json,hashlib,statistics as S,re
from pathlib import Path
import openpyxl
P=Path(__file__).resolve().parent
src=P/'mendeley-supporting-values.xlsx'
assert hashlib.sha256(src.read_bytes()).hexdigest()=='a3eec1dd81423727283939d394e3a51a8c404e2c377a840ba3ccb5946098ec5f'
w=openpyxl.load_workbook(src,data_only=True)
def desc(vals):
    return {'n_source_values':len(vals),'mean':S.mean(vals),'median':S.median(vals),'sd':S.stdev(vals) if len(vals)>1 else None,'min':min(vals),'max':max(vals)}
out={'source':{'doi':'10.17632/3fy2pj3cr6.1','sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'version':'2025-05-28 deposition for Cell Reports Medicine submission; final2026 figure concordance unconfirmed'},'curves':[],'source_auc':[],'fiber_lengths':[],'pATR_positive_nuclei':[],'signature_counts':[],'eligibility':{}}
for s in w:
    strings=[c.value for row in s for c in row if isinstance(c.value,str)]
    hits=[{'cell':c.coordinate,'value':c.value} for row in s for c in row if isinstance(c.value,str) and re.search(r'(?:^|[_-])EMC[123]?|USZ.?23',c.value,re.I)]
    out['eligibility'][s.title]={'description':s['A2'].value,'EMC_labels':hits,'string_labels':list(dict.fromkeys(strings))}
s=w['2E-2L']
for base in range(4,75,10):
    drug=s.cell(base,1).value
    for col in (2,5,8,11,14,17):
        label=s.cell(base+1,col).value
        pts=[]
        for row in range(base+2,base+8):
            vals=[s.cell(row,c).value for c in range(col,col+3)]
            assert all(isinstance(v,(int,float)) for v in vals)
            pts.append({'row':row,'dose_uM':s.cell(row,1).value,'source_cells':[s.cell(row,c).coordinate for c in range(col,col+3)],'values':vals,**desc(vals)})
        out['curves'].append({'sheet':s.title,'drug':drug,'model':label,'units':'source normalized response; final assay duration not recovered','replication':'three source columns, not independent EMC donors','points':pts})
s=w['2M-2O']
for col in (2,5,8,11,14,17):
    for row in range(6,14):
        out['source_auc'].append({'sheet':s.title,'model':s.cell(4,col).value,'drug_label':s.cell(row,1).value,'cells':[s.cell(row,c).coordinate for c in range(col,col+3)],'mean':s.cell(row,col).value,'source_SD':s.cell(row,col+1).value,'source_N':s.cell(row,col+2).value,'warning':'N=13 source field, NOT13donors; calculation units and replication meaning unresolved'})
for sheet in ('4D','4E-F'):
    s=w[sheet]
    for col in range(1,s.max_column+1):
        vals=[s.cell(row,col).value for row in range(6,s.max_row+1) if isinstance(s.cell(row,col).value,(int,float))]
        out['fiber_lengths'].append({'sheet':sheet,'condition':s.cell(5,col).value,'column':s.cell(5,col).column_letter,'values':vals,**desc(vals),'unit':'source fiber length; physical unit/dose/exposure/batch assignment not recovered','replication':'individual fibers, not donors; row order not assumed paired'})
s=w['5B']
for col in range(1,7):
    vals=[s.cell(row,col).value for row in range(7,10)]
    out['pATR_positive_nuclei'].append({'sheet':s.title,'model':s.cell(6,col).value,'values':vals,**desc(vals),'replication':'three source summaries; exact fields/biological replicate identity unconfirmed'})
s=w['2A']
for col in range(2,8):
    out['signature_counts'].append({'sheet':s.title,'model':s.cell(5,col).value,'counts':{s.cell(row,1).value:s.cell(row,col).value for row in range(6,16)},'interpretation':'transcript counts, not protein, HRD state, dependency, or functional knockout'})
out['unresolved_model_sheets']=['4I cleaved caspase3','5E pCHK1 Western blot; need main figure caption to identify model']
(P/'mendeley-analysis.json').write_text(json.dumps(out,indent=2),encoding='utf-8')
print(json.dumps({'all_curves':len(out['curves']),'all_curve_points':sum(len(x['points']) for x in out['curves']),'EMC_curves':[{'drug':x['drug'],'points':[{'dose_uM':y['dose_uM'],'mean':y['mean'],'min':y['min'],'max':y['max']} for y in x['points']]} for x in out['curves'] if 'EMC' in x['model']],'EMC_fibers':[{k:v for k,v in x.items() if k!='values'} for x in out['fiber_lengths'] if x['condition'].startswith('EMC')],'EMC_pATR':[x for x in out['pATR_positive_nuclei'] if 'EMC' in x['model']]},indent=2))
