import json, hashlib, re
from pathlib import Path
import openpyxl
P=Path(__file__).resolve().parent
p=P/'mendeley-supporting-values.xlsx'
assert hashlib.sha256(p.read_bytes()).hexdigest()=='a3eec1dd81423727283939d394e3a51a8c404e2c377a840ba3ccb5946098ec5f'
w=openpyxl.load_workbook(p,data_only=True)
out={}
for s in w:
    rows=[]
    hits=[]
    for row in s:
        cells={c.coordinate:c.value for c in row if c.value is not None}
        if cells: rows.append(cells)
        for c in row:
            if isinstance(c.value,str) and re.search(r'EMC|USZ.?23|time|day|auc|berzos|ceralas|elimus|camon|prexas|rabus|adavos|azenos',c.value,re.I):
                hits.append({'cell':c.coordinate,'value':c.value})
    out[s.title]={'nrows':s.max_row,'ncols':s.max_column,'rows':rows,'hits':hits}
(P/'mendeley-workbook-cells.json').write_text(json.dumps(out,indent=2),encoding='utf-8')
print(json.dumps({k:{'nrows':v['nrows'],'ncols':v['ncols'],'EMC_labels':[h for h in v['hits'] if re.search(r'(?:^|[_-])EMC[123]?|USZ.?23',h['value'],re.I)]} for k,v in out.items()},indent=2))
