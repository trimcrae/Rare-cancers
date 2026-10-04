import pathlib,urllib.request,json,hashlib,io,datetime,warnings,re
import openpyxl
R=pathlib.Path(__file__).resolve().parent
u='https://journals.plos.org/plosone/article/file?type=supplementary&id=10.1371/journal.pone.0201046.s002'
rec={'url':u,'accessed_utc':datetime.datetime.now(datetime.timezone.utc).isoformat()}
try:
 with urllib.request.urlopen(urllib.request.Request(u,headers={'User-Agent':'Mozilla/5.0'}),timeout=35) as f:d=f.read(3*1024*1024+1);rec.update(status=f.status,content_type=f.headers.get('Content-Type'))
 assert len(d)<=3*1024*1024
 (R/'yoo2018-s002.bin').write_bytes(d);rec.update(bytes=len(d),sha256=hashlib.sha256(d).hexdigest())
 with warnings.catch_warnings():
  warnings.simplefilter('ignore');w=openpyxl.load_workbook(io.BytesIO(d),read_only=True,data_only=True)
 rec['sheets']=[]
 for s in w:
  rows=list(s.values)
  rec['sheets'].append({'name':s.title,'rows':s.max_row,'columns':s.max_column,'header':rows[0],'string_alias_hits':[[i+1,str(c)] for i,row in enumerate(rows) for c in row if isinstance(c,str) and re.search(r'\bEMC\b|H.?EMC.?SS|extraskeletal|chondrosarcoma',c,re.I)]})
except Exception as e:rec['error']=str(e)
(R/'citation-gate.json').write_text(json.dumps(rec,indent=2),encoding='utf-8');print(json.dumps(rec,indent=2))
