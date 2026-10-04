"""Scoped imaging eligibility evaluation. No image intensities or outcome selection.
Use installed Python/openpyxl; all received source bytes remain immutable.
"""
from pathlib import Path
from collections import Counter, defaultdict
import hashlib, json, re, xml.etree.ElementTree as ET
from openpyxl import load_workbook
P=Path(__file__).resolve().parent
text=lambda e:' '.join(' '.join(e.itertext()).split())
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
receipts=json.loads((P/'receipts.json').read_text(encoding='utf-8-sig'))
verified=[]
for r in receipts:
    if r.get('file') and r.get('sha256'):
        f=P/r['file']; assert f.stat().st_size==r['bytes']; assert sha(f)==r['sha256'];verified.append(r['file'])
wb=load_workbook(P/'soft-clinical.xlsx',read_only=True,data_only=True)
rows=list(wb['Clinical Information'].values)
soft=[{'id':r[0],'histology':r[3],'MSKCC_type':r[4]} for r in rows[1:] if isinstance(r[0],str) and r[0].startswith('STS_')]
assert len(soft)==51 and len({r['id'] for r in soft})==51
aliases=re.compile(r'extra.?skeletal.*myxoid.*chondro|\bEMCS?\b|NR4A[23]|\bTEC\b|\bCHN\b',re.I)
soft_hits=[r for r in soft if aliases.search(str(r))]
assert not soft_hits
# Evaluate all series and unique patients, preserving absence of diagnosis fields.
rows=list(load_workbook(P/'qin-digest.xlsx',read_only=True,data_only=True).active.values)
headers=rows[0]; data=[dict(zip(headers,r)) for r in rows[1:] if r[0]]
assert len(data)==2168
patients=defaultdict(list)
for r in data: patients[r['Patient ID']].append(r)
assert len(patients)==15
qin=[]
for pid,series in sorted(patients.items()):
    studies=defaultdict(list)
    for r in series: studies[r['Study Instance UID']].append(r)
    qin.append({'patient_id':pid,'series_count':len(series),'study_count':len(studies),
      'sex':sorted({r['Patient Sex'] for r in series if r['Patient Sex']}),
      'age_as_released':sorted({r['Patient Age'] for r in series if r['Patient Age']}),
      'admitting_diagnosis':sorted({r['Admitting Diagnosis Description'] for r in series if r['Admitting Diagnosis Description']}),
      'studies':[{'date_shifted_as_released':s[0]['Study Date'],'series':len(s),'longitudinal_event_labels':sorted({r['Longitudinal Temporal Event Type'] for r in s if r['Longitudinal Temporal Event Type']})} for s in studies.values()],
      'histology_authentication':'Unresolved: released manifest lacks individual histology and trial/paper crosswalk.'})
assert sum(r['study_count'] for r in qin)==38
assert all(not r['admitting_diagnosis'] for r in qin)
root=ET.parse(P/'qin2016.xml').getroot()
t=next(t for t in root.findall('.//table-wrap') if 'Histologic Tumor Subtype' in text(t))
individual=[[text(c) for c in r if c.tag in ['td','th']] for r in t.findall('.//tr')]
individual=[r for r in individual if r and r[0].isdigit()]
assert len(individual)==20
q2016=[{'source_patient_number':r[0],'histology':r[3]} for r in individual]
assert not any(aliases.search(str(r)) for r in q2016)
# Extract every EMC-labelled main table row and every figure caption, without
# assigning pooled results or unlabelled example images to an individual EMC.
explicit={}
for stem in ['myxoid2025','integration2024','multimodal2024']:
    root=ET.parse(P/(stem+'.xml')).getroot()
    emc=[[text(c) for c in r if c.tag in ['td','th']] for r in root.findall('.//tr') if 'chondrosarcoma' in text(r).lower()]
    figs=[{'id':f.get('id'),'caption':text(f)} for f in root.findall('.//fig')]
    explicit[stem]={'emc_table_rows':emc,'figures':figs,'individual_crosswalk':'Not supplied in inspected article/supplements'}
assert 'not yet approved by GEO curators' in (P/'GSE262937-unapproved.html').read_text()
# Reuse the verified CPTAC case evaluation; do not copy or reanalyze protein values.
old=Path(r'C:/Users/mcrae/.codex/worktrees/emc-fresh-20261004-functional/research/autonomy/fresh-discovery-2026-10-04-round3/proteomics')
man=json.loads((old/'MANIFEST.json').read_text())['files']
reuse={}
for name in ['cptac-sar-case-evaluation.json','cptac-eligibility-evaluation.json']:
    expected=next(r for r in man if r['path']==name)
    assert sha(old/name)==expected['sha256']
    reuse[name]={'absolute_path':str(old/name),'sha256':expected['sha256']}
cases=json.loads((old/'cptac-sar-case-evaluation.json').read_text())
assert len(cases)==88 and sum(c['slide_rows'] for c in cases)==305
assert not any(aliases.search(' '.join(c['histologies'])) for c in cases)
result={'source_files_verified':verified,'soft_all51_histology_rows':soft,
 'soft_histology_counts':dict(Counter(r['histology'] for r in soft)),
 'soft_emc_alias_hits':soft_hits,'qin_all15_patient_conditions':qin,
 'qin2016_all20_individual_histologies':q2016,'explicit_EMC_published_cohorts':explicit,
 'GSE262937_status':'Public accession viewer says not yet approved; no data accessed',
 'CPTAC_verified_evaluation_reused':reuse,
 'decision':'SHELVE scoped imaging inference: no authenticated accessible individual EMC imaging contrast; not global absence or a biological negative result.'}
(P/'eligibility-evaluation.json').write_text(json.dumps(result,indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps({'sources_verified':len(verified),'Soft_tissue_Sarcoma_patients':len(soft),'QIN_series':len(data),'QIN_patients':len(qin),'QIN_studies':sum(r['study_count'] for r in qin),'QIN_2016_histologies':len(q2016),'reused_CPTAC_cases':len(cases),'decision':result['decision']},indent=2))
