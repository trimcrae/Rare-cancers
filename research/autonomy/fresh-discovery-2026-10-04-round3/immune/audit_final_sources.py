"""Reproduce bounded source/eligibility checks, not a clinical efficacy estimator."""
from pathlib import Path
from collections import Counter
import hashlib,json,re,datetime,xml.etree.ElementTree as E,zipfile
import openpyxl
from pypdf import PdfReader
B=Path(__file__).parent
out={'utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'scientific_finding':None}
out['source_magic']={}
for p in B.glob('*'):
 if p.suffix in ['.pdf','.xlsx']:
  prefix=p.read_bytes()[:8]
  out['source_magic'][p.name]={'valid_container':prefix.startswith(b'%PDF') if p.suffix=='.pdf' else prefix.startswith(b'PK'),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
w=openpyxl.load_workbook(B/'subramanian2024-tables.xlsx',read_only=True,data_only=True)
out['subramanian_tables']={}
for sn in ['Table S18','Table S19','Table S21']:
 rows=[list(r) for r in list(w[sn].values)[2:] if any(v is not None for v in r)]
 out['subramanian_tables'][sn]={'histology_label_counts':dict(Counter(r[3] for r in rows if r[3] is not None)),'rows':rows}
out['additional_geo']={}
for fn in ['gse239561.txt','gse214779.txt','gse172043.txt']:
 rs=[]
 for rec in re.split(r'(?=\^SAMPLE = )',(B/fn).read_text(encoding='utf8')):
  if '^SAMPLE' not in rec:continue
  lines=[l for l in rec.splitlines() if l.startswith(('^SAMPLE','!Sample_title','!Sample_characteristics_ch1'))]
  rs.append(lines)
 out['additional_geo'][fn]={'n_records':len(rs),'all_identity_records':rs,'explicit_EMC_records':[r for r in rs if re.search(r'NR4A3|extraskeletal.*myxoid|\bEMC\b',str(r),re.I)]}
out['starzer']={'candidate_cases':[{'center':1,'id':7,'histology':'Extra skeletal myxoid chondrosarcoma','site':'hip','sex':'m','age':54,'drug':'Pembrolizumab','PFS_months':3,'best_iRECIST':'PD','eligible':True,'individual_immune_numbers':None},{'center':1,'id':5,'histology':'Myxoid chondrosarcoma','site_literal':'rips','sex':'m','age':65,'drug':'Pembrolizumab','PFS_months':5.9,'best_iRECIST':'PD','eligible':None,'reason':'Extraskeletal/NR4A3 identity unresolved; do not silently include or exclude.'}], 'all_35_patient_table_scope':'Supplement PDF pages9-12 (20Vienna+15Essen); both candidate cases retained; no additional explicitEMC. Potential Schur2017 overlap not resolved.', 'numeric_immune_outcome_model_run':False}
out['starzer']['supplement_spreadsheets']={s.title:{'rows':s.max_row,'columns':s.max_column,'first_row':list(next(s.values))} for s in openpyxl.load_workbook(B/'PMC7993298-jitc-2020-001458supp002.xlsx',read_only=True,data_only=True)}
r=PdfReader(B/'moura2024.pdf')
out['moura_authentication']={'cohort_EMC':4,'paired_transcript_EMC':3,'paired_blood_EMC':3,'paired_EMC_crosswalk_available':False,'do_not_assume_same_3':True,'culture_identity':(r.pages[14].extract_text() or ''),'data_availability':(r.pages[16].extract_text() or '')}
out['moura_authentication']['paired_fraction_each_assay']=3/4
out['moura_authentication']['fraction_interpretation']='Coverage arithmetic only; does not measure immune change, treatment effect, or a new disease finding.'
out['pollack_followup']={'source':'PMC7489365-jamaoncol-e203689-s002.pdf','supplement_pages':16,'EMC_enrolled':1,'IHC_table_header_n':33,'IHC_assessable_total':29,'patient_IDs_in_IHC_table':False,'PDL1_HScore_zero_n':19,'TIL_zero_n':23,'EMC_individual_values':None,'EMC_response_point':None,'within_patient_delta_estimable':False,'visual_receipts':['pollack-supp-7.png','pollack-supp-14.png','pollack-supp-16.png','pollack2020-figure.jpg'],'interpretation':'Manual source-table transcription independently visually checked. Counts are pan-cohort, not EMC values. Chondro pools3conventionalCS+1clearcellCS+1EMC.'}
out['kelly_followup']={'EMC_enrolled':1,'relevant_file':'Supplement1 jamaoncol-6-402-s001.pdf','protocol_file':'Supplement2 jamaoncol-6-402-s002.pdf','eTable4_evaluated':False,'EMC_individual_values':None,'within_patient_delta_estimable':False,'access_evidence':['kelly-followup-retrieval.json','kelly-supp1-followup.json','kelly-cdn-receipt.json','kelly-oa-api-pmc-receipt.json'],'interpretation':'Material unresolved source. Failed retrieval not a biologicalnegative or proofthat no publiclinkeddataexist.'}
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
out['moura_all_supplement_tables']={}
for p in B.glob('moura-*.docx'):
 with zipfile.ZipFile(p) as z:root=E.fromstring(z.read('word/document.xml'))
 out['moura_all_supplement_tables'][p.name]=[[[' '.join(''.join(c.itertext()).split()) for c in row.findall('w:tc',ns)] for row in t.findall('w:tr',ns)] for t in root.findall('.//w:tbl',ns)]
(B/'source-evaluation.json').write_text(json.dumps(out,indent=2,ensure_ascii=False),encoding='utf8')
print(json.dumps({'source_magic_failures':[k for k,v in out['source_magic'].items() if not v['valid_container']],'additional_geo_counts':{k:v['n_records'] for k,v in out['additional_geo'].items()},'subramanian_histologies':{k:v['histology_label_counts'] for k,v in out['subramanian_tables'].items()},'moura_pair_coverage':.75,'scientific_finding':None},indent=2))
